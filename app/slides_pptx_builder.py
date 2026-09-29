"""Build editable PPTX slides from the analyzed slide catalog."""

from __future__ import annotations

import io
import re
from pathlib import Path
from typing import Any

from pptx import Presentation
from pptx.chart.data import ChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION, XL_LEGEND_POSITION, XL_TICK_LABEL_POSITION
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Pt

from .font_registry import collect_embeddable_fonts
from .pptx_fill_export import apply_catalog_fill
from .pptx_font_embed import embed_fonts_in_pptx
from .pptx_paragraph_bullets import apply_paragraph_bullet, apply_paragraph_margins
from .pptx_stroke_export import apply_line_arrow_ends
from .pptx_text_runs_export import add_styled_runs_paragraph, group_text_runs_by_paragraph
from .slide_layer_filter import filter_slide_render_layers
from .slides_export_queue import (
    build_export_queue,
    export_queue_sort_key,
    flex_stack_uses_compact_display_layout,
    paragraph_spacing_for_flex_stack_item,
    resolve_flex_stack_body_insets_pt,
    resolve_text_group_line_gap_pt,
    should_apply_paragraph_spacing_in_flex_stack,
    stack_geometry_pt,
)

EMU_PER_PT = 12700
_ALIGNMENTS = {
    "l": PP_ALIGN.LEFT,
    "ctr": PP_ALIGN.CENTER,
    "r": PP_ALIGN.RIGHT,
    "just": PP_ALIGN.JUSTIFY,
}
_PATH_TOKEN_RE = re.compile(r"([MLCQAZ])|(-?\d*\.?\d+(?:e[-+]?\d+)?)", re.IGNORECASE)
_FILL_MASK_SHAPE_MAP = {
    "roundRect": MSO_SHAPE.ROUNDED_RECTANGLE,
    "round1Rect": MSO_SHAPE.ROUND_1_RECTANGLE,
    "round2SameRect": MSO_SHAPE.ROUND_2_SAME_RECTANGLE,
    "round2DiagRect": MSO_SHAPE.ROUND_2_DIAG_RECTANGLE,
    "ellipse": MSO_SHAPE.OVAL,
}


def _apply_fill_mask_adjustments(shape, mask: dict | None) -> None:
    if not mask:
        return
    adjustments = shape.adjustments
    if mask.get("adj1") is not None and len(adjustments) > 0:
        adjustments[0] = min(max(float(mask["adj1"]), 0.0), 0.5)
    if mask.get("adj2") is not None and len(adjustments) > 1:
        adjustments[1] = min(max(float(mask["adj2"]), 0.0), 0.5)


def pt_to_emu(value: float | int | None) -> int:
    return int(round(float(value or 0) * EMU_PER_PT))


def geometry_box(element: dict, geometry_key: str = "geometry_pt") -> tuple[int, int, int, int]:
    geometry = element.get(geometry_key) or element.get("geometry_pt") or {}
    return (
        pt_to_emu(geometry.get("x_pt")),
        pt_to_emu(geometry.get("y_pt")),
        pt_to_emu(geometry.get("width_pt")),
        pt_to_emu(geometry.get("height_pt")),
    )


def apply_catalog_transform(shape, element: dict) -> None:
    rotate = element.get("rotate") or {}
    if rotate.get("deg"):
        shape.rotation = float(rotate["deg"])

    flip = element.get("flip") or {}
    if not flip.get("h") and not flip.get("v"):
        return

    xfrm = shape.element.spPr.xfrm
    if flip.get("h"):
        xfrm.flipH = True
    if flip.get("v"):
        xfrm.flipV = True


def disable_shape_effects(shape) -> None:
    """Blank layout inherits Office theme shadows; exported shapes should stay flat."""
    try:
        shape.shadow.inherit = False
    except (AttributeError, NotImplementedError):
        # GraphicFrame (table/chart) has no shadow API in python-pptx yet.
        pass


def parse_hex_color(value: str | None) -> RGBColor | None:
    if not value or not isinstance(value, str):
        return None
    cleaned = value.strip().lstrip("#")
    if len(cleaned) == 6:
        return RGBColor(int(cleaned[0:2], 16), int(cleaned[2:4], 16), int(cleaned[4:6], 16))
    return None


def resolve_color(value: Any) -> RGBColor | None:
    if isinstance(value, str):
        return parse_hex_color(value)
    if isinstance(value, dict):
        return parse_hex_color(value.get("color"))
    return None


def apply_stroke(shape, stroke: dict | None) -> None:
    if not stroke:
        shape.line.fill.background()
        return
    line = shape.line
    line.width = Pt(stroke.get("width_pt") or 0.75)
    color = resolve_color(stroke.get("color"))
    if color is not None:
        line.color.rgb = color
    dash = stroke.get("dash")
    if dash and dash != "solid":
        line.dash_style = 2 if dash == "dash" else 1
    apply_line_arrow_ends(shape, stroke)


def apply_solid_fill(shape, fill: dict | None) -> None:
    apply_catalog_fill(shape, fill)


def _ensure_run_typeface(r_pr, tag: str, family: str) -> None:
    node = r_pr.find(qn(f"a:{tag}"))
    if node is None:
        node = OxmlElement(f"a:{tag}")
        r_pr.append(node)
    node.set("typeface", family)


def apply_typography(run, typography: dict | None) -> None:
    typo = typography or {}
    family = typo.get("family")
    if family:
        run.font.name = family
    size_pt = typo.get("size_pt")
    if size_pt is not None:
        run.font.size = Pt(size_pt)
    run.font.bold = bool(typo.get("bold"))
    run.font.italic = bool(typo.get("italic"))
    run.font.underline = bool(typo.get("underline") and typo.get("underline") != "none")
    color = resolve_color(typo.get("color"))
    if color is not None:
        run.font.color.rgb = color
    # LibreOffice's PPTX -> PDF renderer can split Cyrillic words by characters
    # when text runs have no language/proofing metadata. The browser preview
    # forbids arbitrary word breaks; carry that intent into DrawingML runs.
    r_pr = run._r.get_or_add_rPr()
    if family:
        _ensure_run_typeface(r_pr, "cs", family)
        _ensure_run_typeface(r_pr, "ea", family)
    r_pr.set("lang", str(typo.get("lang") or "ru-RU"))
    r_pr.set("dirty", "0")
    r_pr.set("noProof", "1")


def apply_typography_to_paragraph_runs(paragraph, typography: dict | None) -> None:
    if not paragraph.runs:
        apply_typography(paragraph.add_run(), typography)
        return
    for run in paragraph.runs:
        apply_typography(run, typography)


def apply_paragraph_style(paragraph, typography: dict | None, spacing: dict | None) -> None:
    typo = typography or {}
    spacing = spacing or {}
    alignment = _ALIGNMENTS.get(typo.get("alignment") or spacing.get("alignment"))
    if alignment is not None:
        paragraph.alignment = alignment
    level = spacing.get("level")
    if isinstance(level, int):
        paragraph.level = max(level, 0)
    space_before = spacing.get("space_before") or spacing.get("space_before_pt")
    if space_before is not None:
        paragraph.space_before = Pt(space_before)
    space_after = spacing.get("space_after") or spacing.get("space_after_pt")
    if space_after is not None:
        paragraph.space_after = Pt(space_after)
    ratio = spacing.get("line_spacing_ratio") or typo.get("line_height_ratio")
    if ratio is not None:
        paragraph.line_spacing = ratio
    apply_paragraph_margins(paragraph, spacing)


def _resolve_metric_typographies(
    element: dict,
    typography: dict | None = None,
) -> tuple[dict, dict, dict]:
    metric = element.get("metric") or {}
    base = typography or element.get("typography") or {}
    value_typography = {**base, **(metric.get("value_typography") or {})}
    unit_typography = {**base, **(metric.get("unit_typography") or {})}
    return metric, value_typography, unit_typography


def _add_metric_paragraph(
    text_frame,
    *,
    index: int,
    element: dict,
    typography: dict | None,
    spacing: dict | None,
    bullet: dict | None = None,
) -> None:
    metric, value_typography, unit_typography = _resolve_metric_typographies(element, typography)
    paragraph = text_frame.paragraphs[0] if index == 0 else text_frame.add_paragraph()
    value_run = paragraph.runs[0] if paragraph.runs else paragraph.add_run()
    value_run.text = metric.get("value") or ""
    apply_typography(value_run, value_typography)
    unit_run = paragraph.add_run()
    unit_run.text = metric.get("unit") or ""
    apply_typography(unit_run, unit_typography)
    apply_paragraph_style(paragraph, typography or element.get("typography"), spacing)
    apply_paragraph_bullet(paragraph, bullet)


def _add_text_paragraph(
    text_frame,
    *,
    index: int,
    text: str,
    typography: dict | None,
    spacing: dict | None,
    bullet: dict | None = None,
    metric_element: dict | None = None,
) -> None:
    if metric_element and metric_element.get("metric"):
        _add_metric_paragraph(
            text_frame,
            index=index,
            element=metric_element,
            typography=typography,
            spacing=spacing,
            bullet=bullet,
        )
        return
    paragraph = text_frame.paragraphs[0] if index == 0 else text_frame.add_paragraph()
    paragraph.text = text or ""
    apply_paragraph_style(paragraph, typography, spacing)
    apply_paragraph_bullet(paragraph, bullet)
    apply_typography_to_paragraph_runs(paragraph, typography)


def apply_text_frame_margins(text_frame, body_insets: dict | None) -> None:
    insets = body_insets or {}
    text_frame.margin_left = Pt(insets.get("left") or insets.get("left_pt") or 0)
    text_frame.margin_right = Pt(insets.get("right") or insets.get("right_pt") or 0)
    text_frame.margin_top = Pt(insets.get("top") or insets.get("top_pt") or 0)
    text_frame.margin_bottom = Pt(insets.get("bottom") or insets.get("bottom_pt") or 0)


def apply_vertical_anchor(text_frame, anchor: str | None) -> None:
    mapping = {
        "t": MSO_ANCHOR.TOP,
        "ctr": MSO_ANCHOR.MIDDLE,
        "b": MSO_ANCHOR.BOTTOM,
    }
    text_frame.vertical_anchor = mapping.get(anchor or "t", MSO_ANCHOR.TOP)


def add_textbox(slide, element: dict) -> None:
    left, top, width, height = geometry_box(element)
    if width <= 0 or height <= 0:
        return
    shape = slide.shapes.add_textbox(left, top, width, height)
    disable_shape_effects(shape)
    text_frame = shape.text_frame
    wrap = element.get("wrap")
    text_frame.word_wrap = wrap not in ("none", None) if wrap is not None else True
    apply_text_frame_margins(text_frame, element.get("body_insets_pt"))
    apply_vertical_anchor(text_frame, element.get("vertical_anchor"))

    paragraphs = element.get("text_paragraphs") or []
    text = element.get("text") or ""
    if not paragraphs:
        paragraphs = [{"text": text}]

    run_groups = group_text_runs_by_paragraph(element.get("text_runs"))
    use_styled_runs = bool(run_groups) and not element.get("metric")

    for index, paragraph_data in enumerate(paragraphs):
        spacing = paragraph_data.get("paragraph_spacing_pt") or element.get("paragraph_spacing_pt")
        bullet = paragraph_data.get("bullet") or element.get("bullet")
        if element.get("metric") and index == 0:
            _add_metric_paragraph(
                text_frame,
                index=index,
                element=element,
                typography=element.get("typography"),
                spacing=spacing,
                bullet=bullet,
            )
            continue
        if use_styled_runs and index < len(run_groups) and run_groups[index]:
            add_styled_runs_paragraph(
                text_frame,
                index=index,
                runs=run_groups[index],
                typography=element.get("typography"),
                spacing=spacing,
                bullet=bullet,
                apply_paragraph_style=apply_paragraph_style,
                apply_paragraph_bullet=apply_paragraph_bullet,
                apply_typography=apply_typography,
            )
            continue
        _add_text_paragraph(
            text_frame,
            index=index,
            text=paragraph_data.get("text") or "",
            typography=element.get("typography"),
            spacing=spacing,
            bullet=bullet,
        )

    apply_catalog_transform(shape, element)


def add_flex_stack_textbox(slide, members: list[dict], layout: str = "column") -> None:
    if not members:
        return
    if layout == "row":
        for member in members:
            add_textbox(slide, member)
        return
    anchor = members[0]
    geometry = stack_geometry_pt(members)
    left, top, width, height = geometry_box({"geometry_pt": geometry})
    if width <= 0 or height <= 0:
        return

    compact_display_layout = flex_stack_uses_compact_display_layout(members)
    line_gap_pt = resolve_text_group_line_gap_pt(members)
    group_spacing = anchor.get("text_group_spacing_pt") or {}

    shape = slide.shapes.add_textbox(left, top, width, height)
    disable_shape_effects(shape)
    text_frame = shape.text_frame
    text_frame.word_wrap = True
    apply_text_frame_margins(text_frame, resolve_flex_stack_body_insets_pt(members))
    apply_vertical_anchor(text_frame, anchor.get("vertical_anchor"))

    for index, member in enumerate(members):
        spacing = {}
        if should_apply_paragraph_spacing_in_flex_stack(
            member,
            compact_display_layout=compact_display_layout,
        ):
            spacing = paragraph_spacing_for_flex_stack_item(member, index)
        elif compact_display_layout:
            ratio = group_spacing.get("render_line_height_ratio")
            if ratio is not None:
                spacing = {"line_spacing_ratio": ratio}

        if index > 0 and line_gap_pt > 0:
            existing_before = spacing.get("space_before") or spacing.get("space_before_pt") or 0
            spacing = {**spacing, "space_before": float(existing_before) + line_gap_pt}

        typography = member.get("typography") or anchor.get("typography")
        _add_text_paragraph(
            text_frame,
            index=index,
            text=member.get("text") or "",
            typography=typography,
            spacing=spacing or member.get("paragraph_spacing_pt"),
            bullet=member.get("bullet"),
            metric_element=member,
        )

    apply_catalog_transform(shape, anchor)


def add_image(slide, element: dict, assets_dir: Path) -> None:
    from .pptx_image_export import export_image

    export_image(slide, element, assets_dir, disable_shape_effects=disable_shape_effects)


def add_fill_shape(slide, element: dict) -> None:
    left, top, width, height = geometry_box(element)
    if width <= 0 or height <= 0:
        return
    mask = element.get("mask") or {}
    outline_path = mask.get("path") if mask.get("kind") == "path" else None

    if outline_path:
        points = _path_points(outline_path, width, height, left, top)
        if len(points) >= 2:
            builder = slide.shapes.build_freeform(points[0][0], points[0][1])
            builder.add_line_segments(points[1:], close=outline_path.strip().upper().endswith("Z"))
            shape = builder.convert_to_shape()
            disable_shape_effects(shape)
            apply_catalog_fill(shape, element.get("fill"))
            apply_stroke(shape, element.get("stroke"))
            apply_catalog_transform(shape, element)
            return

    if mask.get("kind") in _FILL_MASK_SHAPE_MAP:
        shape = slide.shapes.add_shape(_FILL_MASK_SHAPE_MAP[mask["kind"]], left, top, width, height)
        _apply_fill_mask_adjustments(shape, mask)
    elif mask.get("kind") == "ellipse":
        shape = slide.shapes.add_shape(MSO_SHAPE.OVAL, left, top, width, height)
    else:
        shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, height)
    disable_shape_effects(shape)
    apply_catalog_fill(shape, element.get("fill"))
    apply_stroke(shape, element.get("stroke"))
    apply_catalog_transform(shape, element)


def _cubic_point(p0, p1, p2, p3, t: float) -> tuple[float, float]:
    u = 1 - t
    x = u**3 * p0[0] + 3 * u**2 * t * p1[0] + 3 * u * t**2 * p2[0] + t**3 * p3[0]
    y = u**3 * p0[1] + 3 * u**2 * t * p1[1] + 3 * u * t**2 * p2[1] + t**3 * p3[1]
    return x, y


def _path_points(path_data: str, width: int, height: int, origin_x: int, origin_y: int) -> list[tuple[int, int]]:
    tokens = _PATH_TOKEN_RE.findall(path_data)
    flat: list[str] = []
    for command, number in tokens:
        if command:
            flat.append(command.upper())
        elif number:
            flat.append(number)

    points: list[tuple[int, int]] = []
    index = 0
    current = (0.0, 0.0)
    start = (0.0, 0.0)

    def read_pair() -> tuple[float, float]:
        nonlocal index
        x = float(flat[index])
        y = float(flat[index + 1])
        index += 2
        return x, y

    def append_point(x: float, y: float) -> None:
        points.append((
            origin_x + int(round(x * width)),
            origin_y + int(round(y * height)),
        ))

    while index < len(flat):
        command = flat[index]
        index += 1
        if command == "M":
            current = read_pair()
            start = current
            append_point(*current)
        elif command == "L":
            current = read_pair()
            append_point(*current)
        elif command == "C":
            c1 = read_pair()
            c2 = read_pair()
            end = read_pair()
            for step in range(1, 9):
                t = step / 8
                point = _cubic_point(current, c1, c2, end, t)
                append_point(*point)
            current = end
        elif command == "Z":
            current = start
            append_point(*current)
    return points


def add_line_shape(slide, element: dict) -> None:
    left, top, width, height = geometry_box(element)
    if width <= 0 and height <= 0:
        return
    stroke = element.get("stroke")
    line = element.get("line") or {}
    outline_path = element.get("outline_path")

    if outline_path:
        points = _path_points(outline_path, width, height, left, top)
        if len(points) >= 2:
            builder = slide.shapes.build_freeform(points[0][0], points[0][1])
            builder.add_line_segments(points[1:], close=outline_path.strip().upper().endswith("Z"))
            shape = builder.convert_to_shape()
            disable_shape_effects(shape)
            shape.fill.background()
            apply_stroke(shape, stroke)
            return

    if line.get("outline") == "rect":
        shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, height)
        disable_shape_effects(shape)
        shape.fill.background()
        apply_stroke(shape, stroke)
        return

    x1 = left + int(round(width * float(line.get("x1", 0))))
    y1 = top + int(round(height * float(line.get("y1", 0))))
    x2 = left + int(round(width * float(line.get("x2", 1))))
    y2 = top + int(round(height * float(line.get("y2", 1))))
    connector = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, x1, y1, x2, y2)
    disable_shape_effects(connector)
    apply_stroke(connector, stroke)


def add_table_shape(slide, element: dict) -> None:
    from .pptx_table_export import export_table_shape

    export_table_shape(slide, element, disable_shape_effects=disable_shape_effects)


def _fill_variant_color(variant: Any) -> RGBColor | None:
    if not isinstance(variant, dict):
        return resolve_color(variant)
    stops = variant.get("stops") or []
    if variant.get("kind") in {None, "solid"} or len(stops) < 2:
        return resolve_color(variant.get("color"))
    return resolve_color((stops[0] or {}).get("color"))


def _chart_paint_colors(style: dict) -> list[RGBColor]:
    """Same order the DOM preview paints: fill variant, then series palette."""
    variants = style.get("series_fill_variants") or []
    palette = style.get("series_palette") or []
    if variants:
        colors: list[RGBColor] = []
        for index in range(max(len(variants), len(palette))):
            color = _fill_variant_color(variants[index]) if index < len(variants) else None
            if color is None and index < len(palette):
                color = resolve_color(palette[index])
            if color is not None:
                colors.append(color)
        return colors
    return [color for item in palette if (color := resolve_color(item)) is not None]


def _paint_chart_point(point, color: RGBColor) -> None:
    point.format.fill.solid()
    point.format.fill.fore_color.rgb = color
    point.format.line.fill.solid()
    point.format.line.color.rgb = color


def _numeric_chart_value(value: Any) -> float:
    try:
        return float(value)
    except (TypeError, ValueError):
        return 0.0


def _chart_type(chart_type: str):
    return {
        "line": XL_CHART_TYPE.LINE,
        "area": XL_CHART_TYPE.AREA,
        "pie": XL_CHART_TYPE.PIE,
        "doughnut": XL_CHART_TYPE.DOUGHNUT,
        "bar": XL_CHART_TYPE.COLUMN_CLUSTERED,
    }.get(str(chart_type or "bar").lower(), XL_CHART_TYPE.COLUMN_CLUSTERED)


def _apply_chart_font(font, typography: dict | None) -> None:
    typography = typography or {}
    font.name = typography.get("family") or "Arial"
    font.size = Pt(typography.get("size_pt") or 10)
    font.bold = bool(typography.get("bold"))
    color = resolve_color(typography.get("color") or "#595959")
    if color is not None:
        font.color.rgb = color


def _apply_chart_axis(axis, axis_style: dict, axis_line: dict, grid_line: dict, *, grid_key: str) -> None:
    visible = axis_style.get("visible", True)
    axis.tick_label_position = (
        XL_TICK_LABEL_POSITION.NEXT_TO_AXIS if visible else XL_TICK_LABEL_POSITION.NONE
    )
    _apply_chart_font(axis.tick_labels.font, axis_style.get("typography"))

    if axis_line.get("visible"):
        axis_color = resolve_color(axis_line.get("color") or "#D9D9D9")
        axis.format.line.width = Pt(axis_line.get("width_pt") or 0.75)
        if axis_color is not None:
            axis.format.line.color.rgb = axis_color
    else:
        axis.format.line.fill.background()

    show_grid = bool(grid_line.get("visible") and grid_line.get(grid_key))
    axis.has_major_gridlines = show_grid
    if show_grid:
        grid_color = resolve_color(grid_line.get("color") or "#E6E6E6")
        axis.major_gridlines.format.line.width = Pt(grid_line.get("width_pt") or 0.5)
        if grid_color is not None:
            axis.major_gridlines.format.line.color.rgb = grid_color


def _apply_chart_data_labels(chart, chart_type: str, data_labels: dict, legend: dict) -> None:
    """Value labels (bars/points) or pie percentages, in the planned DS font.

    The frontend label plan (chart-labels.js) decides visibility, size (<= body)
    and a colour that contrasts with the surface under the chart.
    """
    if not chart.plots:
        return
    plot = chart.plots[0]
    circular = chart_type in {"pie", "doughnut", "donut"}
    if circular:
        # Pie: percentages outside the slices, on the slide surface, next to
        # the legend with category names. Doughnut labels would sit on the
        # coloured ring, so the legend alone carries the categories.
        if chart_type != "pie" or not legend.get("visible"):
            return
        typography = legend.get("typography") or {}
        plot.has_data_labels = True
        labels = plot.data_labels
        labels.show_percentage = True
        labels.show_value = False
        labels.show_category_name = False
        labels.number_format = "0%"
        labels.number_format_is_linked = False
        labels.position = XL_LABEL_POSITION.OUTSIDE_END
        _apply_chart_font(labels.font, typography)
        return
    if not data_labels.get("visible"):
        return
    plot.has_data_labels = True
    labels = plot.data_labels
    labels.show_value = True
    labels.show_percentage = False
    labels.show_category_name = False
    if chart_type == "bar":
        labels.position = XL_LABEL_POSITION.OUTSIDE_END
    elif chart_type == "line":
        labels.position = XL_LABEL_POSITION.ABOVE
    _apply_chart_font(labels.font, data_labels.get("typography"))


def add_chart_shape(slide, element: dict) -> None:
    left, top, width, height = geometry_box(element)
    if width <= 0 or height <= 0:
        return
    payload = element.get("chart") or {}
    categories = payload.get("categories_preview") or element.get("categories_preview") or []
    series = payload.get("series") or element.get("series_preview") or []
    if not series:
        series = [{"name": "Series 1", "values_preview": [0 for _ in categories]}]
    point_count = max([len(categories), *[len(item.get("values_preview") or item.get("values") or []) for item in series]], default=1)
    categories = [str(value) for value in categories]
    if len(categories) < point_count:
        categories.extend(str(index + 1) for index in range(len(categories), point_count))

    chart_data = ChartData()
    chart_data.categories = categories
    for index, item in enumerate(series):
        values = list(item.get("values_preview") or item.get("values") or [])
        values.extend([0] * max(0, point_count - len(values)))
        chart_data.add_series(
            str(item.get("name") or item.get("label") or f"Series {index + 1}"),
            [_numeric_chart_value(value) for value in values[:point_count]],
        )

    style = payload.get("style_tokens") or element.get("style_tokens") or {}
    legend = style.get("legend") or {}
    label_plan = payload.get("label_plan") or {}
    chart_type = str(element.get("chart_type") or payload.get("type") or "bar").lower()
    circular_chart = chart_type in {"pie", "doughnut", "donut"}
    cartesian_chart = not circular_chart
    legend_plan = label_plan.get("legend") or {}
    legend_visible = bool(
        legend_plan.get("visible")
        if "visible" in legend_plan
        else (len(series) > 1 if cartesian_chart else legend.get("visible", len(series) > 1))
    )
    chart_left, chart_top, chart_width, chart_height = left, top, width, height
    graphic_frame = slide.shapes.add_chart(
        _chart_type(element.get("chart_type") or payload.get("type")),
        chart_left,
        chart_top,
        chart_width,
        chart_height,
        chart_data,
    )
    chart = graphic_frame.chart
    chart.has_title = False
    chart.has_legend = legend_visible
    if chart.has_legend:
        right = legend_plan.get("position") == "right"
        chart.legend.position = XL_LEGEND_POSITION.RIGHT if right else XL_LEGEND_POSITION.BOTTOM
        chart.legend.include_in_layout = True
        _apply_chart_font(chart.legend.font, legend.get("typography") or {})

    _apply_chart_data_labels(chart, chart_type, style.get("data_labels") or {}, legend)
    if chart_type in {"doughnut", "donut"} and chart.plots:
        subtype = payload.get("subtype") or element.get("subtype") or {}
        hole_size = int(subtype.get("hole_size") or 68)
        hole_size_nodes = chart._chartSpace.xpath(".//c:doughnutChart/c:holeSize")
        if hole_size_nodes:
            hole_size_nodes[0].set("val", str(max(10, min(90, hole_size))))
    if chart_type not in {"pie", "doughnut", "donut"}:
        axis_line = style.get("axis_line") or {"visible": False}
        grid_line = style.get("grid_line") or {"visible": False}
        _apply_chart_axis(
            chart.category_axis,
            style.get("category_axis") or {"visible": True},
            axis_line,
            grid_line,
            grid_key="vertical_visible",
        )
        _apply_chart_axis(
            chart.value_axis,
            style.get("value_axis") or {"visible": True},
            axis_line,
            grid_line,
            grid_key="horizontal_visible",
        )

    palette = _chart_paint_colors(style)
    per_point = chart_type in {"pie", "doughnut", "donut"} or (chart_type == "bar" and len(series) == 1)
    if palette and chart.plots:
        # PowerPoint treats a val-less <c:varyColors/> as false and then ignores per-point srgb fills.
        plot = chart.plots[0]
        plot.vary_by_categories = per_point
        vary_colors = plot._element.varyColors
        if vary_colors is not None:
            vary_colors.set("val", "1" if per_point else "0")
    if per_point:
        for plotted_series in chart.series:
            for index, point in enumerate(plotted_series.points):
                if not palette:
                    continue
                _paint_chart_point(point, palette[index % len(palette)])
    else:
        for index, plotted_series in enumerate(chart.series):
            if not palette:
                continue
            color = palette[index % len(palette)]
            plotted_series.format.fill.solid()
            plotted_series.format.fill.fore_color.rgb = color
            plotted_series.format.line.fill.solid()
            plotted_series.format.line.color.rgb = color

    if chart_type in {"doughnut", "donut"}:
        circular_layout = style.get("circular_layout") or {}
        center_style = circular_layout.get("center_metric") or {}
        center_metric = (
            element.get("center_metric_preview")
            or payload.get("center_metric_preview")
            or center_style.get("preview")
        )
        if center_metric and center_style.get("enabled", True):
            hole_size = int((payload.get("subtype") or element.get("subtype") or {}).get("hole_size") or 68)
            # Keep the same safe inner diameter as the web preview. The remaining
            # part of the hole is a visual gap between the label and the ring.
            free_diameter = int(width * max(0.30, min(0.82 * hole_size / 100, 0.74)))
            metric_height = min(free_diameter, pt_to_emu(
                max(
                    float((center_style.get("value_typography") or {}).get("size_pt") or 18),
                    float((center_style.get("unit_typography") or {}).get("size_pt") or 12),
                ) * 1.45
            ))
            metric_box = slide.shapes.add_textbox(
                left + (width - free_diameter) // 2,
                top + (width - metric_height) // 2,
                free_diameter,
                metric_height,
            )
            disable_shape_effects(metric_box)
            frame = metric_box.text_frame
            frame.clear()
            frame.word_wrap = False
            frame.margin_left = frame.margin_right = frame.margin_top = frame.margin_bottom = 0
            frame.vertical_anchor = MSO_ANCHOR.MIDDLE
            paragraph = frame.paragraphs[0]
            paragraph.alignment = PP_ALIGN.CENTER
            value_run = paragraph.add_run()
            value_run.text = str(center_metric.get("value") or "")
            apply_typography(value_run, center_style.get("value_typography") or {})
            unit_run = paragraph.add_run()
            unit_run.text = str(center_metric.get("unit") or "")
            apply_typography(unit_run, center_style.get("unit_typography") or {})


def _diagram_style(element: dict) -> tuple[dict, dict, dict]:
    style = (element.get("diagram") or {}).get("style_tokens") or element.get("style_tokens") or {}
    connector = dict(style.get("connector") or {})
    arrow = style.get("arrow") or {}
    connector["head"] = connector.get("head") or {"type": arrow.get("head_type") or "none"}
    connector["tail"] = connector.get("tail") or {"type": arrow.get("tail_type") or "triangle"}
    return connector, style.get("node") or {}, (style.get("node") or {}).get("typography") or {}


def _diagram_layout(width: int, height: int, count: int, vertical: bool) -> tuple[int, int, int]:
    """Return node_width, node_height, gap that fit inside the diagram box."""
    count = max(count, 1)
    gaps = max(count - 1, 0)
    if vertical:
        preferred_gap = pt_to_emu(14 if count <= 3 else 10 if count <= 5 else 6)
        max_gap = max(pt_to_emu(4), height // max(count * 5, 1))
        gap = min(preferred_gap, max_gap) if gaps else 0
        node_height = max(pt_to_emu(22), (height - gap * gaps) // count)
        if node_height * count + gap * gaps > height and gaps:
            gap = max(pt_to_emu(3), (height - node_height * count) // gaps)
            if gap < 0:
                gap = 0
                node_height = max(pt_to_emu(18), height // count)
        node_width = max(pt_to_emu(48), int(width * (0.78 if count <= 4 else 0.88)))
        node_width = min(node_width, width)
        return node_width, node_height, gap

    preferred_gap = pt_to_emu(16 if count <= 3 else 10 if count <= 5 else 6)
    max_gap = max(pt_to_emu(4), width // max(count * 5, 1))
    gap = min(preferred_gap, max_gap) if gaps else 0
    node_width = max(pt_to_emu(28), (width - gap * gaps) // count)
    if node_width * count + gap * gaps > width and gaps:
        gap = max(pt_to_emu(3), (width - node_width * count) // gaps)
        if gap < 0:
            gap = 0
            node_width = max(pt_to_emu(22), width // count)
    node_height = max(pt_to_emu(28), int(height * (0.42 if count <= 4 else 0.5)))
    node_height = min(node_height, height)
    return node_width, node_height, gap


def add_diagram_shape(slide, element: dict) -> None:
    left, top, width, height = geometry_box(element)
    if width <= 0 or height <= 0:
        return
    diagram = element.get("diagram") or {}
    labels = diagram.get("preview_texts") or element.get("preview_texts") or ["Step 1", "Step 2", "Step 3"]
    labels = [str(label) for label in labels[:8]]
    if not labels:
        return
    connector_style, node_style, typography = _diagram_style(element)
    vertical = str(element.get("diagram_type") or diagram.get("diagram_type") or "flow").lower() == "process"
    count = len(labels)
    node_width, node_height, gap = _diagram_layout(width, height, count, vertical)
    if vertical:
        node_left = left + max(0, int((width - node_width) / 2))
    else:
        node_top = top + max(0, int((height - node_height) / 2))

    font_size = float(typography.get("size_pt") or 12)
    if count >= 7:
        font_size = min(font_size, 9)
    elif count >= 5:
        font_size = min(font_size, 10)

    arrow_size = "sm" if count >= 5 else "med"
    connector_stroke = {
        **connector_style,
        "head": {
            **(connector_style.get("head") or {"type": "none"}),
            "width": (connector_style.get("head") or {}).get("width") or arrow_size,
            "length": (connector_style.get("head") or {}).get("length") or arrow_size,
        },
        "tail": {
            **(connector_style.get("tail") or {"type": "triangle"}),
            "width": (connector_style.get("tail") or {}).get("width") or arrow_size,
            "length": (connector_style.get("tail") or {}).get("length") or arrow_size,
        },
    }

    nodes = []
    for index, label in enumerate(labels):
        if vertical:
            node_top = top + index * (node_height + gap)
        else:
            node_left = left + index * (node_width + gap)
        radius_pt = float(node_style.get("radius_pt") if node_style.get("radius_pt") is not None else 8)
        radius_ratio = float(node_style.get("radius_ratio") or 0)
        if radius_pt > 0 or radius_ratio > 0:
            shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, node_left, node_top, node_width, node_height)
            short_side_pt = min(node_width, node_height) / 12700
            adjustment = radius_ratio if radius_ratio > 0 else radius_pt / max(short_side_pt, 0.01)
            if shape.adjustments:
                shape.adjustments[0] = max(0.0, min(0.5, adjustment))
        else:
            shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, node_left, node_top, node_width, node_height)
        disable_shape_effects(shape)
        apply_solid_fill(shape, node_style.get("fill"))
        border = node_style.get("border") or {
            "color": connector_style.get("color"),
            "width_pt": connector_style.get("width_pt") or 1,
        }
        apply_stroke(shape, border)
        frame = shape.text_frame
        frame.clear()
        frame.word_wrap = True
        margin = 4 if count >= 5 else 6
        frame.margin_left = frame.margin_right = Pt(margin)
        frame.margin_top = frame.margin_bottom = Pt(3 if count >= 5 else 4)
        paragraph = frame.paragraphs[0]
        paragraph.alignment = _ALIGNMENTS.get(typography.get("alignment"), PP_ALIGN.CENTER)
        run = paragraph.add_run()
        run.text = label
        apply_typography(run, typography)
        run.font.size = Pt(font_size)
        run.font.bold = bool(typography.get("bold"))
        text_color = resolve_color(typography.get("color"))
        if text_color is not None:
            run.font.color.rgb = text_color
        nodes.append(shape)

    for start, end in zip(nodes, nodes[1:]):
        if vertical:
            x1 = start.left + int(start.width / 2)
            y1 = start.top + start.height
            x2 = end.left + int(end.width / 2)
            y2 = end.top
        else:
            x1 = start.left + start.width
            y1 = start.top + int(start.height / 2)
            x2 = end.left
            y2 = end.top + int(end.height / 2)
        line = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, x1, y1, x2, y2)
        disable_shape_effects(line)
        apply_stroke(line, connector_stroke)


def export_queue_item(slide, item: dict, assets_dir: Path) -> None:
    if item.get("kind") == "flex_stack":
        add_flex_stack_textbox(slide, item.get("members") or [], item.get("layout") or "column")
        return
    export_element(slide, item.get("element") or {}, assets_dir)


def export_element(slide, element: dict, assets_dir: Path) -> None:
    kind = element.get("kind")
    if kind == "text":
        add_textbox(slide, element)
    elif kind == "image":
        add_image(slide, element, assets_dir)
    elif kind == "fill":
        add_fill_shape(slide, element)
    elif kind == "line":
        add_line_shape(slide, element)
    elif kind == "table":
        add_table_shape(slide, element)
    elif kind == "chart":
        add_chart_shape(slide, element)
    elif kind == "diagram":
        add_diagram_shape(slide, element)


def slide_export_items(slide: dict) -> list[dict]:
    layers = filter_slide_render_layers((slide.get("render") or {}).get("layers"), slide)
    layer_items = [{"kind": "element", "element": layer} for layer in layers]
    content_queue = build_export_queue(slide.get("content_elements") or [])
    return sorted([*layer_items, *content_queue], key=export_queue_sort_key)


def build_editable_pptx_from_catalog(report: dict, assets_dir: Path) -> bytes:
    catalog = report.get("slides") or {}
    slides = sorted(catalog.get("slides") or [], key=lambda item: item.get("slide_number") or 0)
    if not slides:
        raise ValueError("Slide catalog is empty")

    first_render = slides[0].get("render") or {}
    slide_size = first_render.get("slide_size_pt") or {"width": 960, "height": 540}
    prs = Presentation()
    prs.slide_width = pt_to_emu(slide_size.get("width"))
    prs.slide_height = pt_to_emu(slide_size.get("height"))
    blank_layout = prs.slide_layouts[6]

    for slide_data in slides:
        slide = prs.slides.add_slide(blank_layout)
        background = (slide_data.get("render") or {}).get("background_color")
        if background:
            fill = slide.background.fill
            fill.solid()
            color = parse_hex_color(background)
            if color is not None:
                fill.fore_color.rgb = color
        for item in slide_export_items(slide_data):
            export_queue_item(slide, item, assets_dir)

    output = io.BytesIO()
    prs.save(output)
    pptx_bytes = output.getvalue()
    fonts_root = Path(__file__).resolve().parent / "fonts"
    embeddable = collect_embeddable_fonts(report, fonts_root)
    if embeddable:
        pptx_bytes = embed_fonts_in_pptx(pptx_bytes, embeddable)
    return pptx_bytes
