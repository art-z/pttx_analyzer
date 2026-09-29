"""Build a PPTX deck from rendered slide images."""

from __future__ import annotations

import base64
import io
import re

from pptx import Presentation

EMU_PER_PT = 12700
_DATA_URL_RE = re.compile(r"^data:(?P<mime>[^;]+);base64,(?P<payload>.+)$", re.DOTALL)


def decode_data_url(data_url: str) -> bytes:
    match = _DATA_URL_RE.match(data_url.strip())
    if not match:
        raise ValueError("Expected a base64 data URL image payload")
    return base64.b64decode(match.group("payload"))


def build_pptx_from_slide_images(
    images: list[str],
    width_pt: float,
    height_pt: float,
) -> bytes:
    if not images:
        raise ValueError("At least one slide image is required")

    prs = Presentation()
    prs.slide_width = int(round(width_pt * EMU_PER_PT))
    prs.slide_height = int(round(height_pt * EMU_PER_PT))
    blank_layout = prs.slide_layouts[6]

    for data_url in images:
        slide = prs.slides.add_slide(blank_layout)
        image_bytes = decode_data_url(data_url)
        slide.shapes.add_picture(
            io.BytesIO(image_bytes),
            0,
            0,
            width=prs.slide_width,
            height=prs.slide_height,
        )

    output = io.BytesIO()
    prs.save(output)
    return output.getvalue()
