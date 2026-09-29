"""Export catalog tables to editable PPTX with DOM-aligned styling."""

from __future__ import annotations

from typing import Any

from pptx.enum.text import MSO_ANCHOR
from pptx.oxml import parse_xml
from pptx.oxml.ns import nsdecls, qn
from pptx.util import Pt

from .pptx_fill_export import build_solid_color_fill_xml
from .table_styles import (
    INVISIBLE_TABLE_BORDER,
    STYLE_PART_MERGE_ORDER,
    cell_style_part_names,
    is_invisible_table_border,
    merge_cell_styles,
)

EMU_PER_PT = 12700
_BORDER_SIDE_TAGS = {
    "left": "lnL",
    "right": "lnR",
    "top": "lnT",
    "bottom": "lnB",
}
_VERTICAL_ANCHORS = {
    "t": MSO_ANCHOR.TOP,
    "ctr": MSO_ANCHOR.MIDDLE,
    "b": MSO_ANCHOR.BOTTOM,
}


def pt_to_emu(value: float | int | None) -> int:
    return int(round(float(value or 0) * EMU_PER_PT))


def column_weights_for_count(weights: list[float] | None, count: int) -> list[float]:
    if count <= 0:
        return [1.0]
    if not weights:
        return [1.0] * count
    if count <= len(weights):
        return list(weights[:count])
    average = sum(weights) / len(weights) if weights else 1.0
    return [*weights, *([average or 1.0] * (count - len(weights)))]


def table_dimensions(element: dict) -> tuple[int, int, list[list[str]]]:
    table_meta = element.get("table") or {}
    preview = element.get("preview") or table_meta.get("data_preview", {}).get("rows") or []
    row_count = int(
        element.get("rows")
        or table_meta.get("content_model", {}).get("rows")
        or len(table_meta.get("row_heights_pt") or [])
        or len(preview)
        or 1
    )
    col_count = int(
        element.get("cols")
        or table_meta.get("content_model", {}).get("columns")
        or len(table_meta.get("column_widths_pt") or [])
        or (len(preview[0]) if preview else 1)
    )
    matrix = [
        [
            str((preview[row_index][col_index] if row_index < len(preview) and col_index < len(preview[row_index]) else "") or "")
            for col_index in range(col_count)
        ]
        for row_index in range(row_count)
    ]
    return row_count, col_count, matrix


def table_intrinsic_height_pt(table_meta: dict, row_count: int) -> float | None:
    row_heights = table_meta.get("row_heights_pt") or []
    if not row_heights or row_count <= 0:
        return None
    return round(sum(column_weights_for_count(row_heights, row_count)), 2)


def resolve_table_design_height_pt(element: dict) -> float | None:
    table_meta = element.get("table") or {}
    layout_height = table_meta.get("layout_height_pt")
    if layout_height and layout_height > 0:
        return float(layout_height)
    geometry = element.get("geometry_pt") or {}
    shape_height = geometry.get("height_pt") or table_meta.get("max_height_pt")
    if shape_height and shape_height > 0:
        return float(shape_height)
    return None


def is_fixed_row_sizing(table_meta: dict, row_count: int) -> bool:
    sizing_mode = (table_meta.get("structure") or {}).get("row_sizing", {}).get("mode")
    if sizing_mode == "fit_content":
        return False
    if sizing_mode == "fixed":
        return True
    frame_height = table_meta.get("frame_height_pt")
    grid_sum = table_intrinsic_height_pt(table_meta, row_count)
    if frame_height and grid_sum and frame_height > grid_sum + 1:
        return False
    heights = table_meta.get("row_heights_pt") or []
    if len(heights) < 2:
        return False
    rounded = {round(float(value or 0), 1) for value in heights}
    return len(rounded) == 1


def should_stretch_table_height(element: dict, row_count: int) -> bool:
    table_meta = element.get("table") or {}
    sizing_mode = (table_meta.get("structure") or {}).get("row_sizing", {}).get("mode")
    if sizing_mode == "fit_content":
        return bool(resolve_table_design_height_pt(element) and table_meta.get("row_heights_pt"))
    if sizing_mode == "fixed":
        design_height = resolve_table_design_height_pt(element)
        if not design_height:
            return False
        intrinsic = table_intrinsic_height_pt(table_meta, row_count)
        if intrinsic is None:
            return True
        return intrinsic < design_height - 0.5

    frame_height = table_meta.get("frame_height_pt") or (element.get("geometry_pt") or {}).get("height_pt")
    grid_sum = table_intrinsic_height_pt(table_meta, row_count)
    if frame_height and grid_sum and frame_height > grid_sum + 1:
        return bool(table_meta.get("row_heights_pt"))

    design_height = resolve_table_design_height_pt(element)
    if not design_height:
        return False
    intrinsic = table_intrinsic_height_pt(table_meta, row_count)
    if intrinsic is None:
        return True
    return intrinsic < design_height - 0.5


def uses_explicit_row_heights(table_meta: dict, row_count: int) -> bool:
    return is_fixed_row_sizing(table_meta, row_count) and bool(table_meta.get("row_heights_pt")) and row_count > 0


def table_structure_flags(table_meta: dict, row_count: int, col_count: int) -> dict[str, bool]:
    structure = table_meta.get("structure") or {}
    flags = dict(structure.get("flags") or {})
    rules = table_meta.get("style_rules") or {}
    tokens = table_meta.get("style_tokens") or {}

    if structure.get("header_row") is not None:
        flags["first_row"] = True
    if rules.get("firstRow") or tokens.get("first_row_cell") or tokens.get("header_cell"):
        flags["first_row"] = flags.get("first_row", True)
    if rules.get("firstCol") or tokens.get("first_col_cell"):
        flags["first_col"] = True
    if rules.get("lastCol") or tokens.get("last_col_cell"):
        flags["last_col"] = True
    if rules.get("lastRow") or tokens.get("last_row_cell"):
        flags["last_row"] = True
    if any(rules.get(key) or tokens.get(key) for key in ("band1H", "band2H", "band1_row_cell", "band2_row_cell")):
        flags["band_row"] = True
    if any(rules.get(key) or tokens.get(key) for key in ("band1V", "band2V", "band1_col_cell", "band2_col_cell")):
        flags["band_col"] = True
    if flags.get("last_row") is None and row_count > 0:
        flags["last_row"] = False
    if flags.get("last_col") is None and col_count > 0:
        flags["last_col"] = False
    return flags


def resolve_cell_style_from_rules(table_meta: dict, row_index: int, col_index: int, row_count: int, col_count: int) -> dict | None:
    rules = table_meta.get("style_rules") or {}
    if not rules:
        return None
    flags = table_structure_flags(table_meta, row_count, col_count)
    applicable = set(cell_style_part_names(flags, row_index, col_index, row_count, col_count))
    merged = None
    for part_name in STYLE_PART_MERGE_ORDER:
        if part_name not in applicable:
            continue
        part_style = rules.get(part_name)
        if part_style:
            merged = merge_cell_styles(merged, part_style)
    return merged if merged and merged.keys() else None


def resolve_cell_style_legacy(table_meta: dict, row_index: int, col_index: int, row_count: int, col_count: int) -> dict | None:
    tokens = table_meta.get("style_tokens") or {}
    flags = table_structure_flags(table_meta, row_count, col_count)
    header_row = (table_meta.get("structure") or {}).get("header_row")

    if header_row is not None and row_index == header_row:
        return tokens.get("header_cell") or tokens.get("first_row_cell") or tokens.get("body_cell")
    if flags.get("first_col") and col_index == 0:
        return tokens.get("first_col_cell") or tokens.get("body_cell") or tokens.get("whole_cell")
    if flags.get("last_col") and col_index == col_count - 1:
        return tokens.get("last_col_cell") or tokens.get("body_cell") or tokens.get("whole_cell")
    if flags.get("band_row"):
        band_start = 1 if flags.get("first_row") else 0
        if row_index >= band_start:
            ordinal = row_index - band_start
            token = tokens.get("band2_row_cell") if ordinal % 2 else tokens.get("band1_row_cell")
            return token or tokens.get("body_cell") or tokens.get("whole_cell")
    return tokens.get("body_cell") or tokens.get("whole_cell") or tokens.get("header_cell")


def _style_anchor_col(table_meta: dict, preview_row: list[str], row_index: int, col_index: int, has_direct: bool) -> int | None:
    merged_cells = (table_meta.get("structure") or {}).get("merged_cells") or []
    for merge in merged_cells:
        row = merge.get("row", 0)
        col = merge.get("col", 0)
        row_span = merge.get("row_span", 1)
        col_span = merge.get("col_span", 1)
        if row_index < row or row_index >= row + row_span:
            continue
        if col_index < col or col_index >= col + col_span:
            continue
        if row_index == row and col_index == col:
            return None
        return col
    if has_direct or col_index <= 0:
        return None
    if col_index < len(preview_row) and str(preview_row[col_index] or "").strip():
        return None
    for prev_col in range(col_index - 1, -1, -1):
        if prev_col < len(preview_row) and str(preview_row[prev_col] or "").strip():
            return prev_col
    return None


def resolve_cell_style(
    element: dict,
    row_index: int,
    col_index: int,
    row_count: int,
    col_count: int,
    preview_row: list[str],
) -> dict | None:
    table_meta = element.get("table") or {}
    style_grid = table_meta.get("cell_styles") or []
    direct = None
    if row_index < len(style_grid) and col_index < len(style_grid[row_index]):
        direct = style_grid[row_index][col_index]
    from_rules = resolve_cell_style_from_rules(table_meta, row_index, col_index, row_count, col_count)
    legacy = resolve_cell_style_legacy(table_meta, row_index, col_index, row_count, col_count)

    if direct:
        merged = direct
    elif from_rules:
        merged = from_rules
    else:
        merged = legacy

    anchor_col = _style_anchor_col(table_meta, preview_row, row_index, col_index, direct is not None)
    if anchor_col is not None and anchor_col != col_index:
        return resolve_cell_style(element, row_index, anchor_col, row_count, col_count, preview_row)
    return merged


def resolve_table_cell_edge_borders(
    style_grid: list[list[dict | None]],
    row_index: int,
    col_index: int,
    row_count: int,
    col_count: int,
    *,
    last_col: int | None = None,
) -> dict[str, Any]:
    current = (style_grid[row_index][col_index] or {}).get("borders") or {} if style_grid else {}
    above = (
        (style_grid[row_index - 1][col_index] or {}).get("borders") or {}
        if row_index > 0 and style_grid
        else {}
    )
    left_cell = (
        (style_grid[row_index][col_index - 1] or {}).get("borders") or {}
        if col_index > 0 and style_grid
        else {}
    )
    edge_last_col = last_col if last_col is not None else col_index
    last_cell = (
        (style_grid[row_index][edge_last_col] or {}).get("borders") or current
        if style_grid and edge_last_col < len(style_grid[row_index])
        else current
    )
    borders: dict[str, Any] = {}
    borders["top"] = current.get("top") if row_index == 0 else (above.get("bottom") or current.get("top"))
    borders["left"] = current.get("left") if col_index == 0 else (left_cell.get("right") or current.get("left"))
    borders["right"] = last_cell.get("right") if edge_last_col == col_count - 1 else None
    borders["bottom"] = current.get("bottom") if row_index == row_count - 1 else None
    return {key: value for key, value in borders.items() if value}


def _is_outer_table_edge(side: str, row_index: int, col_index: int, row_count: int, col_count: int) -> bool:
    if side == "left":
        return col_index == 0
    if side == "right":
        return col_index == col_count - 1
    if side == "top":
        return row_index == 0
    if side == "bottom":
        return row_index == row_count - 1
    return False


def apply_cell_border_side(tc_pr, side: str, stroke: dict | None) -> None:
    if not stroke:
        return
    tag = _BORDER_SIDE_TAGS.get(side)
    if not tag:
        return
    existing = tc_pr.find(qn(f"a:{tag}"))
    if existing is not None:
        tc_pr.remove(existing)

    if is_invisible_table_border(stroke):
        width_emu = pt_to_emu(stroke.get("width_pt") or 0.75)
        line = parse_xml(
            f'<a:{tag} {nsdecls("a")} w="{width_emu}">'
            f'<a:solidFill><a:srgbClr val="000000"><a:alpha val="0"/></a:srgbClr></a:solidFill>'
            f"</a:{tag}>"
        )
        tc_pr.append(line)
        return

    fill_xml = build_solid_color_fill_xml(stroke.get("color"))
    if not fill_xml:
        return

    width_emu = pt_to_emu(stroke.get("width_pt") or 0.75)
    dash = stroke.get("dash")
    dash_xml = ""
    if dash and dash != "solid":
        dash_xml = f'<a:prstDash val="{dash}"/>'
    line = parse_xml(
        f'<a:{tag} {nsdecls("a")} w="{width_emu}">'
        f"{fill_xml}"
        f"{dash_xml}"
        f"</a:{tag}>"
    )
    tc_pr.append(line)


def resolve_export_cell_borders(
    cell_style: dict | None,
    row_index: int,
    col_index: int,
    row_count: int,
    col_count: int,
) -> dict[str, Any]:
    """Write per-cell tcPr borders like the source PPTX (not edge-deduped preview borders)."""
    borders = dict((cell_style or {}).get("borders") or {})
    for side in ("left", "right", "top", "bottom"):
        if not _is_outer_table_edge(side, row_index, col_index, row_count, col_count):
            continue
        stroke = borders.get(side)
        if not stroke or is_invisible_table_border(stroke):
            borders[side] = dict(INVISIBLE_TABLE_BORDER)
    return borders


def apply_cell_borders(cell, borders: dict | None) -> None:
    if not borders:
        return
    tc_pr = cell._tc.get_or_add_tcPr()
    for side in ("left", "right", "top", "bottom"):
        stroke = borders.get(side)
        if stroke:
            apply_cell_border_side(tc_pr, side, stroke)


def apply_cell_fill(cell, fill: dict | None) -> None:
    from .pptx_fill_export import apply_catalog_fill

    apply_catalog_fill(cell, fill)


def apply_cell_padding(cell, padding_pt: dict | None) -> None:
    if not padding_pt:
        return
    if padding_pt.get("left") is not None:
        cell.margin_left = Pt(padding_pt["left"])
    if padding_pt.get("right") is not None:
        cell.margin_right = Pt(padding_pt["right"])
    if padding_pt.get("top") is not None:
        cell.margin_top = Pt(padding_pt["top"])
    if padding_pt.get("bottom") is not None:
        cell.margin_bottom = Pt(padding_pt["bottom"])


def apply_cell_vertical_anchor(cell, anchor: str | None) -> None:
    mapped = _VERTICAL_ANCHORS.get(anchor or "t")
    if mapped is not None:
        cell.vertical_anchor = mapped


def _merge_cell_typography(default_typography: dict | None, cell_typography: dict | None) -> dict:
    merged = dict(default_typography or {})
    for key, value in (cell_typography or {}).items():
        if value is not None:
            merged[key] = value
    return merged


def _default_table_text_color(table_meta: dict) -> str:
    tokens = table_meta.get("style_tokens") or {}
    for key in ("body_cell", "whole_cell", "header_cell"):
        color = ((tokens.get(key) or {}).get("typography") or {}).get("color")
        if color:
            return str(color)
    return "#000000"


def resolve_table_cell_typography(
    cell_style: dict | None,
    cell_text: dict | None,
    table_meta: dict,
) -> dict:
    typography = _merge_cell_typography(
        (cell_style or {}).get("typography"),
        (cell_text or {}).get("typography") if cell_text else None,
    )
    if typography.get("color"):
        return typography
    if (cell_style or {}).get("fill", {}).get("color"):
        return typography
    return {**typography, "color": _default_table_text_color(table_meta)}


def _apply_typography_to_paragraph_runs(paragraph, typography: dict | None) -> None:
    from .slides_pptx_builder import apply_typography_to_paragraph_runs

    apply_typography_to_paragraph_runs(paragraph, typography)


def _populate_cell_text(cell, cell_text: dict, typography: dict | None) -> None:
    from .pptx_paragraph_bullets import apply_paragraph_bullet, apply_paragraph_margins
    from .slides_pptx_builder import apply_paragraph_style, apply_typography

    text_frame = cell.text_frame
    text_frame.word_wrap = True
    text_frame.clear()

    segments = cell_text.get("text_segments") or []
    typography = dict(typography or {})

    if segments:
        for index, segment in enumerate(segments):
            paragraph = text_frame.paragraphs[0] if index == 0 else text_frame.add_paragraph()
            paragraph.text = segment.get("text") or ""
            segment_typography = _merge_cell_typography(typography, segment.get("typography"))
            spacing = segment.get("paragraph_spacing_pt")
            apply_paragraph_style(paragraph, segment_typography, spacing)
            apply_paragraph_margins(paragraph, spacing)
            apply_paragraph_bullet(paragraph, segment.get("bullet"))
            _apply_typography_to_paragraph_runs(paragraph, segment_typography)
        return

    text_paragraphs = cell_text.get("text_paragraphs") or []
    if text_paragraphs:
        for index, paragraph_data in enumerate(text_paragraphs):
            paragraph = text_frame.paragraphs[0] if index == 0 else text_frame.add_paragraph()
            paragraph.text = paragraph_data.get("text") or ""
            paragraph_typography = _merge_cell_typography(typography, paragraph_data.get("typography"))
            spacing = paragraph_data.get("paragraph_spacing_pt") or cell_text.get("paragraph_spacing_pt")
            apply_paragraph_style(paragraph, paragraph_typography, spacing)
            apply_paragraph_margins(paragraph, spacing)
            apply_paragraph_bullet(paragraph, paragraph_data.get("bullet") or cell_text.get("bullet"))
            _apply_typography_to_paragraph_runs(paragraph, paragraph_typography)
        return

    paragraph = text_frame.paragraphs[0]
    paragraph.text = cell_text.get("text") or ""
    spacing = cell_text.get("paragraph_spacing_pt")
    apply_paragraph_style(paragraph, typography, spacing)
    apply_paragraph_margins(paragraph, spacing)
    apply_paragraph_bullet(paragraph, cell_text.get("bullet"))
    _apply_typography_to_paragraph_runs(paragraph, typography)


def compute_export_row_heights_pt(element: dict, row_count: int, design_height_pt: float | None) -> list[float]:
    table_meta = element.get("table") or {}
    weights = column_weights_for_count(table_meta.get("row_heights_pt") or [], row_count)
    stretch = should_stretch_table_height(element, row_count)

    if uses_explicit_row_heights(table_meta, row_count):
        intrinsic = sum(weights)
        if stretch and design_height_pt and intrinsic > 0 and intrinsic < design_height_pt - 0.5:
            scale = design_height_pt / intrinsic
            return [round(weight * scale, 2) for weight in weights]
        return weights

    if design_height_pt and row_count > 0:
        if stretch or not weights or sum(weights) <= 0:
            return [round(design_height_pt / row_count, 2)] * row_count
        intrinsic = sum(weights)
        if intrinsic > 0:
            scale = design_height_pt / intrinsic
            return [round(weight * scale, 2) for weight in weights]

    if weights and sum(weights) > 0:
        return weights
    return [20.0] * max(row_count, 1)


def compute_export_column_widths_pt(element: dict, col_count: int, target_width_pt: float) -> list[float]:
    table_meta = element.get("table") or {}
    weights = column_weights_for_count(table_meta.get("column_widths_pt") or [], col_count)
    intrinsic = sum(weights)
    if target_width_pt > 0 and intrinsic > 0:
        scale = target_width_pt / intrinsic
        return [round(weight * scale, 2) for weight in weights]
    if target_width_pt > 0 and col_count > 0:
        return [round(target_width_pt / col_count, 2)] * col_count
    return weights


def _find_merge_at(merged_cells: list[dict], row_index: int, col_index: int) -> dict | None:
    for merge in merged_cells:
        row = merge.get("row", 0)
        col = merge.get("col", 0)
        col_span = merge.get("col_span", 1)
        if row_index != row:
            continue
        if col_index < col or col_index >= col + col_span:
            continue
        return {
            "merge": merge,
            "is_anchor": col_index == col,
            "col_span": col_span,
            "last_col": col + col_span - 1,
        }
    return None


def apply_table_flags(table, table_meta: dict) -> None:
    flags = (table_meta.get("structure") or {}).get("flags") or {}
    if table_meta.get("cell_styles"):
        table.first_row = False
        table.last_row = False
        table.first_col = False
        table.last_col = False
        table.horz_banding = False
        table.vert_banding = False
        return
    table.first_row = bool(flags.get("first_row"))
    table.last_row = bool(flags.get("last_row"))
    table.first_col = bool(flags.get("first_col"))
    table.last_col = bool(flags.get("last_col"))
    table.horz_banding = bool(flags.get("band_row"))
    table.vert_banding = bool(flags.get("band_col"))


def export_table_shape(slide, element: dict, *, disable_shape_effects) -> None:
    from .slides_pptx_builder import geometry_box

    row_count, col_count, matrix = table_dimensions(element)
    if row_count <= 0 or col_count <= 0:
        return

    left, top, width, height = geometry_box(element)
    if width <= 0 or height <= 0:
        return

    table_meta = element.get("table") or {}
    design_height_pt = resolve_table_design_height_pt(element) or round(height / EMU_PER_PT, 2)
    design_width_pt = table_meta.get("layout_width_pt") or round(width / EMU_PER_PT, 2)

    shape = slide.shapes.add_table(row_count, col_count, left, top, width, height)
    disable_shape_effects(shape)
    table = shape.table
    apply_table_flags(table, table_meta)

    column_widths_pt = compute_export_column_widths_pt(element, col_count, design_width_pt)
    for col_index, col_width_pt in enumerate(column_widths_pt):
        if col_index < len(table.columns):
            table.columns[col_index].width = pt_to_emu(col_width_pt)

    row_heights_pt = compute_export_row_heights_pt(element, row_count, design_height_pt)
    for row_index, row_height_pt in enumerate(row_heights_pt):
        if row_index < len(table.rows):
            table.rows[row_index].height = pt_to_emu(row_height_pt)

    style_grid = [
        [
            resolve_cell_style(element, row_index, col_index, row_count, col_count, matrix[row_index])
            for col_index in range(col_count)
        ]
        for row_index in range(row_count)
    ]
    cell_text_grid = table_meta.get("cell_text") or []
    merged_cells = (table_meta.get("structure") or {}).get("merged_cells") or []

    for row_index in range(row_count):
        for col_index in range(col_count):
            merge_info = _find_merge_at(merged_cells, row_index, col_index)
            if merge_info and not merge_info["is_anchor"]:
                continue

            cell = table.cell(row_index, col_index)
            cell_style = style_grid[row_index][col_index]
            export_borders = resolve_export_cell_borders(
                cell_style,
                row_index,
                col_index,
                row_count,
                col_count,
            )

            apply_cell_fill(cell, (cell_style or {}).get("fill"))
            apply_cell_borders(cell, export_borders)
            apply_cell_padding(cell, (cell_style or {}).get("padding_pt"))
            apply_cell_vertical_anchor(cell, (cell_style or {}).get("vertical_anchor"))

            cell_text = None
            if row_index < len(cell_text_grid) and col_index < len(cell_text_grid[row_index]):
                cell_text = cell_text_grid[row_index][col_index]

            blank_stubs = (table_meta.get("structure") or {}).get("blank_stub_cells") or []
            if [row_index, col_index] in blank_stubs:
                cell.text_frame.clear()
                continue

            typography = resolve_table_cell_typography(cell_style, cell_text, table_meta)
            if cell_text:
                _populate_cell_text(cell, cell_text, typography)
            else:
                text_frame = cell.text_frame
                text_frame.clear()
                paragraph = text_frame.paragraphs[0]
                paragraph.text = matrix[row_index][col_index] or ""
                from .slides_pptx_builder import apply_paragraph_style

                apply_paragraph_style(paragraph, typography, None)
                _apply_typography_to_paragraph_runs(paragraph, typography)

    for merge in merged_cells:
        row = merge.get("row", 0)
        col = merge.get("col", 0)
        row_span = merge.get("row_span", 1)
        col_span = merge.get("col_span", 1)
        if row_span <= 1 and col_span <= 1:
            continue
        end_row = min(row + row_span - 1, row_count - 1)
        end_col = min(col + col_span - 1, col_count - 1)
        if end_row == row and end_col == col:
            continue
        try:
            table.cell(row, col).merge(table.cell(end_row, end_col))
        except ValueError:
            continue
