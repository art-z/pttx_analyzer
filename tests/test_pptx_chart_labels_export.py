from pptx import Presentation
from pptx.enum.chart import XL_LEGEND_POSITION
from pptx.util import Pt

from app.slides_pptx_builder import add_chart_shape


def _slide():
    prs = Presentation()
    return prs.slides.add_slide(prs.slide_layouts[6])


def _typography(size=10, color="#FFFFFF"):
    return {"family": "Arial", "size_pt": size, "color": color, "bold": False}


def _element(chart_type, style, *, series=None, plan=None):
    return {
        "kind": "chart",
        "chart_type": chart_type,
        "geometry_pt": {"x_pt": 40, "y_pt": 40, "width_pt": 400, "height_pt": 240},
        "chart": {
            "type": chart_type,
            "categories_preview": ["Разработка", "Маркетинг", "Поддержка", "Инфраструктура"],
            "series": series or [{"name": "Доли", "values_preview": [40, 25, 20, 15]}],
            "style_tokens": style,
            "label_plan": plan or {},
        },
    }


def _chart(slide):
    return [shape for shape in slide.shapes if shape.has_chart][0].chart


def test_bar_with_hidden_value_axis_exports_value_labels_in_planned_font():
    slide = _slide()
    add_chart_shape(slide, _element("bar", {
        "legend": {"visible": False, "typography": _typography()},
        "category_axis": {"visible": True, "typography": _typography(10)},
        "value_axis": {"visible": False, "typography": _typography(10)},
        "data_labels": {"visible": True, "show_value": True, "typography": _typography(10, "#FFFFFF")},
    }))
    chart = _chart(slide)
    plot = chart.plots[0]
    assert plot.has_data_labels
    assert plot.data_labels.show_value is True
    assert plot.data_labels.font.size == Pt(10)
    assert str(plot.data_labels.font.color.rgb) == "FFFFFF"
    assert chart.has_legend is False


def test_pie_keeps_native_legend_and_outside_percentages():
    slide = _slide()
    add_chart_shape(slide, _element("pie", {
        "legend": {"visible": True, "typography": _typography(9.5, "#222222")},
        "data_labels": {"visible": False},
    }, plan={"legend": {"visible": True, "position": "right"}}))
    chart = _chart(slide)
    assert chart.has_legend is True
    assert chart.legend.position == XL_LEGEND_POSITION.RIGHT
    assert chart.legend.font.size == Pt(9.5)
    assert str(chart.legend.font.color.rgb) == "222222"
    legend_text = [shape.text for shape in slide.shapes if getattr(shape, "has_text_frame", False)]
    assert not any("Разработка" in text for text in legend_text)
    labels = chart.plots[0].data_labels
    assert labels.show_percentage is True
    assert labels.font.size == Pt(9.5)



def test_doughnut_keeps_native_legend_inside_chart_object():
    slide = _slide()
    add_chart_shape(slide, _element("doughnut", {
        "legend": {"visible": True, "typography": _typography(9.5, "#222222")},
        "data_labels": {"visible": False},
    }, plan={"legend": {"visible": True, "position": "right"}}))
    chart = _chart(slide)
    assert chart.has_legend is True
    assert chart.legend.position == XL_LEGEND_POSITION.RIGHT
    legend_text = [shape.text for shape in slide.shapes if getattr(shape, "has_text_frame", False)]
    assert not any("Разработка" in text for text in legend_text)



def test_multi_series_bar_enables_bottom_native_legend_even_when_template_hidden():
    slide = _slide()
    add_chart_shape(slide, _element("bar", {
        "legend": {"visible": False, "typography": _typography(11, "#123456")},
        "category_axis": {"visible": True, "typography": _typography(9)},
        "value_axis": {"visible": True, "typography": _typography(9)},
        "data_labels": {"visible": False},
    }, series=[
        {"name": "2025", "values_preview": [1, 2, 3, 4]},
        {"name": "2026", "values_preview": [2, 3, 4, 5]},
    ]))
    chart = _chart(slide)
    assert chart.has_legend is True
    assert chart.legend.position == XL_LEGEND_POSITION.BOTTOM
    assert chart.legend.include_in_layout is True
    assert chart.legend.font.size == Pt(11)
    assert str(chart.legend.font.color.rgb) == "123456"

def test_multi_series_line_keeps_native_legend_without_value_labels():
    slide = _slide()
    add_chart_shape(slide, _element("line", {
        "legend": {"visible": True, "typography": _typography(9)},
        "category_axis": {"visible": True, "typography": _typography(9)},
        "value_axis": {"visible": True, "typography": _typography(9)},
        "data_labels": {"visible": False},
    }, series=[
        {"name": "2025", "values_preview": [1, 2, 3, 4]},
        {"name": "2026", "values_preview": [2, 3, 4, 5]},
    ]))
    chart = _chart(slide)
    assert chart.has_legend is True
    assert chart.legend.position == XL_LEGEND_POSITION.BOTTOM
    assert chart.legend.include_in_layout is True
    legend_text = [shape.text for shape in slide.shapes if getattr(shape, "has_text_frame", False)]
    assert not any("2025" in text for text in legend_text)
    assert not any("2026" in text for text in legend_text)
    assert chart.plots[0].has_data_labels is False
