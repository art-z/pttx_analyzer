import json
import re
import unittest
import zipfile
from io import BytesIO
from pathlib import Path

from lxml import etree
from pptx import Presentation

from app.colors import _build_theme_map
from app.graphic_elements import NS, parse_table_element
from app.pptx_table_export import (
    compute_export_row_heights_pt,
    export_table_shape,
    resolve_cell_style,
    resolve_export_cell_borders,
    uses_explicit_row_heights,
)
from app.slides_pptx_builder import build_editable_pptx_from_catalog, disable_shape_effects
from app.table_styles import is_invisible_table_border


def _table_cell_has_background_fill(cell_xml: str, color: str) -> bool:
    match = re.search(r"<a:tcPr[^>]*>.*?</a:tcPr>", cell_xml, re.S)
    if not match:
        return False
    tc_pr = match.group(0)
    if "<a:noFill/>" in tc_pr:
        return False
    return "<a:solidFill>" in tc_pr and color.upper() in tc_pr.upper()


class PptxTableExportTest(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.fixture_root = Path("output/32f36382e98f")
        cls.header_fixture_root = Path("output/6523b92ae85c")
        if not (cls.fixture_root / "report.json").is_file():
            raise unittest.SkipTest("table fixture report missing")
        cls.report = json.loads((cls.fixture_root / "report.json").read_text())
        if (cls.header_fixture_root / "report.json").is_file():
            cls.header_report = json.loads((cls.header_fixture_root / "report.json").read_text())
        else:
            cls.header_report = None

    def _table_element(self, slide_number: int = 38, *, report=None) -> dict:
        report = report or self.report
        slide = next(item for item in report["slides"]["slides"] if item["slide_number"] == slide_number)
        return next(item for item in slide["content_elements"] if item.get("kind") == "table")

    def test_resolve_cell_style_uses_catalog_fill(self):
        element = self._table_element()
        style = resolve_cell_style(element, 0, 0, element["rows"], element["cols"], ["", "", ""])
        self.assertEqual((style or {}).get("fill", {}).get("color"), "#0077FF")

    def test_slide39_header_row_keeps_highlight_only_in_last_cell(self):
        if not self.header_report:
            self.skipTest("header row fixture report missing")
        element = self._table_element(slide_number=39, report=self.header_report)
        header_row = element["table"]["structure"]["header_row"]
        last_col = element["cols"] - 1
        for col_index in range(element["cols"]):
            style = resolve_cell_style(
                element,
                header_row,
                col_index,
                element["rows"],
                element["cols"],
                [""] * element["cols"],
            )
            fill_color = (style or {}).get("fill", {}).get("color")
            text_color = (style or {}).get("typography", {}).get("color")
            bottom_border = (style or {}).get("borders", {}).get("bottom", {}).get("color", {}).get("color")
            if col_index == last_col:
                self.assertEqual(fill_color, "#0077FF")
                self.assertEqual(text_color, "#EBF3F9")
            else:
                self.assertIsNone(fill_color)
                self.assertIsNone(text_color)
            self.assertEqual(bottom_border, "#0077FF")

    def test_export_slide39_header_row_keeps_highlight_only_in_last_cell(self):
        if not self.header_report:
            self.skipTest("header row fixture report missing")
        element = self._table_element(slide_number=39, report=self.header_report)
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        export_table_shape(slide, element, disable_shape_effects=disable_shape_effects)
        output = BytesIO()
        prs.save(output)
        xml = zipfile.ZipFile(BytesIO(output.getvalue())).read("ppt/slides/slide1.xml").decode()
        import re

        header_cells = re.findall(r"<a:tc>.*?</a:tc>", xml, re.S)[: element["cols"]]
        self.assertEqual(len(header_cells), element["cols"])
        for index, cell_xml in enumerate(header_cells):
            has_fill = _table_cell_has_background_fill(cell_xml, "0077FF")
            has_white_text = "EBF3F9" in cell_xml.upper() and "<a:solidFill>" in cell_xml.split("<a:tcPr", 1)[0]
            has_black_text = "000000" in cell_xml.upper() and "<a:solidFill>" in cell_xml.split("<a:tcPr", 1)[0]
            has_bottom_border = "lnB" in cell_xml and (
                "0077FF" in cell_xml.upper() or 'schemeClr val="dk2"' in cell_xml
            )
            if index == element["cols"] - 1:
                self.assertTrue(has_fill)
                self.assertTrue(has_white_text)
            else:
                self.assertFalse(has_fill)
                self.assertFalse(has_white_text)
                self.assertTrue(has_black_text)
            self.assertTrue(has_bottom_border)

    def test_compute_export_row_heights_stretches_to_design_height(self):
        element = self._table_element()
        table_meta = element["table"]
        self.assertTrue(uses_explicit_row_heights(table_meta, element["rows"]))
        design_height = element["geometry_pt"]["height_pt"]
        row_heights = compute_export_row_heights_pt(element, element["rows"], design_height)
        self.assertEqual(len(row_heights), element["rows"])
        self.assertAlmostEqual(sum(row_heights), design_height, delta=1.0)

    def test_export_table_shape_writes_styles_and_row_heights(self):
        element = self._table_element()
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        export_table_shape(slide, element, disable_shape_effects=disable_shape_effects)
        output = BytesIO()
        prs.save(output)
        xml = zipfile.ZipFile(BytesIO(output.getvalue())).read("ppt/slides/slide1.xml").decode()
        self.assertIn("0077FF", xml.upper())
        self.assertIn("lnL", xml)
        self.assertIn(' h="', xml)

    def test_slide14_exports_semi_transparent_cell_borders(self):
        root = Path("output/d8e58f9bb017")
        source = root / "source.pptx"
        if not source.is_file():
            self.skipTest("d8 fixture missing")
        theme_map = _build_theme_map(json.loads((root / "report.json").read_text()).get("theme") or {})
        with zipfile.ZipFile(source) as archive:
            slide_root = etree.fromstring(archive.read("ppt/slides/slide14.xml"))
            table_styles = etree.fromstring(archive.read("ppt/tableStyles.xml"))
        table = slide_root.find(".//a:tbl", NS)
        package = type("Package", (), {
            "exists": lambda _self, part: part == "ppt/tableStyles.xml",
            "xml": lambda _self, part: table_styles,
        })()
        element = parse_table_element(
            table,
            theme_map,
            {"+mn-lt": "Arial"},
            package=package,
        )
        element["geometry_pt"] = {"x_pt": 0, "y_pt": 0, "width_pt": 300, "height_pt": 400}

        style = resolve_cell_style(element, 1, 1, element["rows"], element["cols"], ["", ""])
        border = (style or {}).get("borders", {}).get("top") or {}
        self.assertAlmostEqual(border.get("width_pt"), 0.75)
        self.assertEqual((border.get("color") or {}).get("color"), "#0077FF")
        self.assertAlmostEqual((border.get("color") or {}).get("alpha"), 0.502, places=3)

        header = resolve_cell_style(element, 0, 0, element["rows"], element["cols"], ["", ""])
        self.assertTrue(is_invisible_table_border((header or {}).get("borders", {}).get("left")))
        self.assertTrue(is_invisible_table_border((header or {}).get("borders", {}).get("top")))

        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        export_table_shape(slide, element, disable_shape_effects=disable_shape_effects)
        output = BytesIO()
        prs.save(output)
        slide_xml = zipfile.ZipFile(BytesIO(output.getvalue())).read("ppt/slides/slide1.xml").decode()
        cell0_tcpr = re.search(r"<a:tc>.*?<a:tcPr[^>]*>.*?</a:tcPr>", slide_xml, re.S)
        self.assertIsNotNone(cell0_tcpr)
        cell0 = cell0_tcpr.group(0)
        self.assertRegex(cell0, r"<a:lnL[^>]*>.*?alpha val=\"0\"", re.S)
        self.assertRegex(cell0, r"<a:lnT[^>]*>.*?alpha val=\"0\"", re.S)
        self.assertIn('0077FF', slide_xml.upper())
        self.assertIn('alpha val="50200"', slide_xml)
        self.assertIn('w="9525"', slide_xml)
        self.assertNotIn('schemeClr val="lt1"', slide_xml)
        self.assertNotIn('w="12700"', slide_xml)
        for tc_pr in re.findall(r"<a:tcPr[^>]*>.*?</a:tcPr>", slide_xml, re.S):
            for side in "LRTB":
                match = re.search(rf"<a:ln{side}[^>]*>.*?</a:ln{side}>", tc_pr, re.S)
                if not match:
                    continue
                border_xml = match.group(0).upper()
                self.assertNotIn("FFFFFF", border_xml)
                self.assertNotIn('SCHEMECLR VAL="LT1"', border_xml)

    def test_slide14_catalog_export_suppresses_table_border_style_on_outer_edges(self):
        root = Path("output/ab9e08a8c1f5")
        if not (root / "report.json").is_file():
            self.skipTest("slide 14 catalog fixture missing")
        report = json.loads((root / "report.json").read_text())
        pptx_bytes = build_editable_pptx_from_catalog(report, root / "assets")
        with zipfile.ZipFile(BytesIO(pptx_bytes)) as archive:
            slide_xml = archive.read("ppt/slides/slide14.xml").decode()
        cell0_tcpr = re.search(r"<a:tc>.*?<a:tcPr[^>]*>.*?</a:tcPr>", slide_xml, re.S)
        self.assertIsNotNone(cell0_tcpr)
        cell0 = cell0_tcpr.group(0)
        self.assertRegex(cell0, r"<a:lnL[^>]*>.*?alpha val=\"0\"", re.S)
        self.assertRegex(cell0, r"<a:lnT[^>]*>.*?alpha val=\"0\"", re.S)
        self.assertNotIn('w="12700"', slide_xml)
        cells = re.findall(r"<a:tc>.*?</a:tc>", slide_xml, re.S)
        cell1_tcpr = re.search(r"<a:tcPr[^>]*>.*?</a:tcPr>", cells[1], re.S)
        self.assertIsNotNone(cell1_tcpr)
        cell1 = cell1_tcpr.group(0)
        self.assertRegex(cell1, r"<a:lnB[^>]*>.*?0077FF.*?alpha val=\"50200\"", re.S | re.I)
        self.assertRegex(cell1, r"<a:lnL[^>]*>.*?0077FF.*?alpha val=\"50200\"", re.S | re.I)

    def test_resolve_export_cell_borders_keeps_all_cell_sides(self):
        cell_style = {
            "borders": {
                "left": {"width_pt": 0.75, "color": {"kind": "solid", "color": "#0077FF", "alpha": 0.502}},
                "bottom": {"width_pt": 0.75, "color": {"kind": "solid", "color": "#0077FF", "alpha": 0.502}},
            }
        }
        borders = resolve_export_cell_borders(cell_style, 0, 1, 11, 2)
        self.assertIn("bottom", borders)
        self.assertIn("left", borders)
        self.assertTrue(is_invisible_table_border(borders["top"]))
        self.assertTrue(is_invisible_table_border(borders["right"]))

    def test_export_multiline_cell_text_styles_all_runs(self):
        element = {
            "rows": 1,
            "cols": 1,
            "geometry_pt": {"x_pt": 0, "y_pt": 0, "width_pt": 120, "height_pt": 40},
            "table": {
                "cell_styles": [[{"typography": {"family": "Arial", "size_pt": 9.0, "color": "#000000", "alignment": "l"}}]],
                "cell_text": [[{
                    "text": "Цифровой ассистент\nдля B2B",
                    "typography": {"family": "Arial", "size_pt": 9.0, "color": "#000000", "alignment": "l"},
                }]],
                "structure": {"flags": {}},
                "style_tokens": {"body_cell": {"typography": {"family": "Arial", "size_pt": 9.0}}},
            },
        }

        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        export_table_shape(slide, element, disable_shape_effects=disable_shape_effects)
        output = BytesIO()
        prs.save(output)
        xml = zipfile.ZipFile(BytesIO(output.getvalue())).read("ppt/slides/slide1.xml").decode()
        run_colors = re.findall(
            r"<a:r>.*?<a:rPr[^>]*>.*?<a:solidFill><a:srgbClr val=\"([0-9A-F]+)\"",
            xml,
            re.S | re.I,
        )
        run_sizes = re.findall(r"<a:r>.*?<a:rPr[^>]* sz=\"(\d+)\"", xml, re.S | re.I)
        self.assertGreaterEqual(len(run_colors), 2)
        self.assertTrue(all(color.upper() == "000000" for color in run_colors))
        self.assertTrue(all(size == "900" for size in run_sizes))
        self.assertEqual(len(re.findall(r"<a:br/>", xml)), 1)
        self.assertEqual(len(re.findall(r"<a:p>", xml)), 1)

    def test_build_editable_pptx_preserves_table_styles(self):
        pptx_bytes = build_editable_pptx_from_catalog(self.report, self.fixture_root / "assets")
        with zipfile.ZipFile(BytesIO(pptx_bytes)) as archive:
            slide_xml = archive.read("ppt/slides/slide38.xml").decode()
        self.assertIn("0077FF", slide_xml.upper())
        self.assertIn("a:tbl", slide_xml)
