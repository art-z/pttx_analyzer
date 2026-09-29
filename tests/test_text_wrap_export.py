import json
import unittest
import zipfile
from io import BytesIO
from pathlib import Path

from pptx import Presentation

from app.pptx import PPTXPackage
from app.pptx_text_runs_export import add_styled_runs_paragraph, group_text_runs_by_paragraph
from app.slides_pptx_builder import apply_paragraph_bullet, apply_paragraph_style, apply_typography, build_editable_pptx_from_catalog
from app.template_layers import build_slide_content_layers
from app.text_wrap import expand_wrapped_paragraph_runs, text_wrap_enabled


class TextWrapTest(unittest.TestCase):
    def test_wrap_enabled_for_square(self):
        self.assertTrue(text_wrap_enabled("square"))
        self.assertFalse(text_wrap_enabled("none"))

    def test_wraps_long_monospace_line_by_width(self):
        paragraphs = [{
            "alignment": "l",
            "runs": [{
                "text": "background: url(https://example.com/" + ("x" * 80) + ")",
                "family": "Consolas",
                "size_pt": 16.0,
                "color": "#FFFFFF",
            }],
            "text": "background: url(https://example.com/" + ("x" * 80) + ")",
        }]
        wrapped = expand_wrapped_paragraph_runs(
            paragraphs,
            geometry_pt={"width_pt": 373.12},
            body_insets_pt={"left": 2.81, "right": 2.81},
            wrap="square",
        )
        breaks = [run for run in wrapped[0]["runs"] if run.get("break") == "line"]
        self.assertGreaterEqual(len(breaks), 2)

    def test_skips_proportional_title_text(self):
        paragraphs = [{
            "alignment": "l",
            "runs": [{
                "text": "Пример слайда-разделителя",
                "family": "Play",
                "size_pt": 48.0,
                "color": "#0077FF",
            }],
            "text": "Пример слайда-разделителя",
        }]
        wrapped = expand_wrapped_paragraph_runs(
            paragraphs,
            geometry_pt={"width_pt": 320.0},
            body_insets_pt={"left": 2.81, "right": 2.81},
            wrap="square",
        )
        self.assertEqual(wrapped[0]["runs"], paragraphs[0]["runs"])
        self.assertNotIn("\n", wrapped[0]["text"])

    def test_preserves_existing_soft_break_before_auto_wrap(self):
        paragraphs = [{
            "alignment": "l",
            "runs": [
                {"text": "line-a", "family": "Consolas", "size_pt": 16.0},
                {"break": "line"},
                {"text": "line-b-" + ("z" * 80), "family": "Consolas", "size_pt": 16.0},
            ],
            "text": "line-a\nline-b-zzz",
        }]
        wrapped = expand_wrapped_paragraph_runs(
            paragraphs,
            geometry_pt={"width_pt": 373.12},
            body_insets_pt={"left": 2.81, "right": 2.81},
            wrap="square",
        )
        runs = wrapped[0]["runs"]
        self.assertEqual(runs[1]["break"], "line")
        self.assertGreater(sum(1 for run in runs if run.get("break") == "line"), 1)


class Slide15CodeBlockTest(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.root = Path("output/da49ce6c4194")
        if not (cls.root / "source.pptx").is_file():
            raise unittest.SkipTest("fixture pptx missing")
        cls.theme = json.loads((cls.root / "report.json").read_text())["theme"]
        cls.package = PPTXPackage(cls.root / "source.pptx")

    def _code_element(self) -> dict:
        elements = build_slide_content_layers(
            self.package,
            "ppt/slides/slide15.xml",
            (960.0, 540.0),
            self.theme,
        )
        return next(element for element in elements if element.get("name") == "Google Shape;407;p47")

    def test_slide_29_section_title_is_not_auto_wrapped(self):
        elements = build_slide_content_layers(
            self.package,
            "ppt/slides/slide29.xml",
            (960.0, 540.0),
            self.theme,
        )
        title = next(element for element in elements if element.get("name") == "Google Shape;808;p61")
        self.assertEqual(title["text"], "Пример слайда-разделителя")
        self.assertNotIn("\n", title["text"])

    def test_code_block_exports_wrap_and_text_runs(self):
        element = self._code_element()
        self.assertEqual(element.get("wrap"), "square")
        self.assertIn("text_runs", element)
        url_paragraph = next(
            paragraph for paragraph in element["text_paragraphs"]
            if "background: url" in paragraph["text"]
        )
        self.assertGreaterEqual(url_paragraph["text"].count("\n"), 2)

    def test_code_block_export_preserves_syntax_colors(self):
        element = self._code_element()
        report = {
            "slides": {
                "slides": [{
                    "slide_number": 15,
                    "content_elements": [element],
                    "render": {"slide_size_pt": {"width": 960, "height": 540}, "layers": []},
                }],
            },
        }
        pptx_bytes = build_editable_pptx_from_catalog(report, self.root / "assets")
        with zipfile.ZipFile(BytesIO(pptx_bytes)) as archive:
            slide_xml = archive.read("ppt/slides/slide1.xml").decode()
        for color in ("FFE59C", "C3A3E2", "FFBD92", "A0DFE0"):
            self.assertIn(color, slide_xml)


class PptxTextRunsExportTest(unittest.TestCase):
    def test_group_text_runs_by_paragraph(self):
        runs = [
            {"text": "a", "color": "#111111"},
            {"break": "line"},
            {"text": "b", "color": "#222222"},
            {"break": "paragraph"},
            {"text": "c", "color": "#333333"},
        ]
        groups = group_text_runs_by_paragraph(runs)
        self.assertEqual(len(groups), 2)
        self.assertEqual(len(groups[0]), 3)
        self.assertEqual(groups[1][0]["text"], "c")

    def test_add_styled_runs_paragraph_writes_line_breaks_and_colors(self):
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        text_frame = slide.shapes.add_textbox(0, 0, 914400, 914400).text_frame
        add_styled_runs_paragraph(
            text_frame,
            index=0,
            runs=[
                {"text": "body", "color": "#FFE59C", "family": "Consolas", "size_pt": 16},
                {"text": " {", "color": "#FFFFFF", "family": "Consolas", "size_pt": 16},
            ],
            typography={"family": "Consolas", "size_pt": 16, "color": "#FFFFFF"},
            spacing=None,
            bullet=None,
            apply_paragraph_style=apply_paragraph_style,
            apply_paragraph_bullet=apply_paragraph_bullet,
            apply_typography=apply_typography,
        )
        xml = text_frame._txBody.xml
        self.assertIn("FFE59C", xml)
        self.assertIn("FFFFFF", xml)


if __name__ == "__main__":
    unittest.main()
