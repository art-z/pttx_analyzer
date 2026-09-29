from app.generated_export_scene import build_generated_export_scene
from app.slides_pptx_builder import build_editable_pptx_from_catalog

from io import BytesIO
from pptx import Presentation


def test_generated_scene_preserves_browser_resolved_graphic_styles():
    report = {"typography": {"visibility": {"slide_size_pt": {"width": 960, "height": 540}}}}
    source = {
        "export_scene_version": 1,
        "content_elements": [{
            "kind": "chart",
            "chart": {"style_tokens": {"series_palette": [{"color": "#123456"}]}},
        }],
    }

    scene = build_generated_export_scene(report, "Demo", [source])

    assert scene["version"] == 1
    assert scene["slide_size_pt"] == {"width": 960, "height": 540}
    assert scene["slides"][0]["slide_number"] == 1
    assert scene["slides"][0]["content_elements"][0]["chart"]["style_tokens"]["series_palette"] == [
        {"color": "#123456"},
    ]
    assert source.get("slide_number") is None


def test_generated_scene_still_styles_older_unresolved_clients():
    report = {
        "colors": {"resolved_palette": [{"color": "#AA5500"}]},
        "typography": {"visibility": {"slide_size_pt": {"width": 960, "height": 540}}},
    }
    source = {"content_elements": [{"kind": "chart", "chart": {"style_tokens": {}}}]}

    scene = build_generated_export_scene(report, "Demo", [source])

    assert scene["slides"][0]["export_scene_version"] == 1
    assert scene["slides"][0]["content_elements"][0]["chart"]["style_tokens"]["series_palette"]


def test_pptx_uses_scene_colors_and_keeps_chart_editable(tmp_path):
    report = {"typography": {"visibility": {"slide_size_pt": {"width": 960, "height": 540}}}}
    source = {
        "export_scene_version": 1,
        "content_elements": [{
            "kind": "chart",
            "chart_type": "bar",
            "geometry_pt": {"x_pt": 80, "y_pt": 90, "width_pt": 600, "height_pt": 300},
            "chart": {
                "categories_preview": ["До", "После"],
                "series": [{"name": "Время", "values_preview": [12, 4]}],
                "style_tokens": {
                    "series_palette": [{"color": "#123456"}, {"color": "#654321"}],
                    "series_fill_variants": [
                        {"kind": "solid", "color": "#123456"},
                        {"kind": "solid", "color": "#654321"},
                    ],
                },
            },
        }],
    }
    scene = build_generated_export_scene(report, "Demo", [source])
    pptx = build_editable_pptx_from_catalog({"slides": {"slides": scene["slides"]}}, tmp_path)
    presentation = Presentation(BytesIO(pptx))
    charts = [shape.chart for shape in presentation.slides[0].shapes if shape.has_chart]

    assert len(charts) == 1
    assert len(charts[0].series[0].points) == 2
    assert str(charts[0].series[0].points[0].format.fill.fore_color.rgb) == "123456"
    assert str(charts[0].series[0].points[1].format.fill.fore_color.rgb) == "654321"


def test_pptx_table_preserves_intentionally_blank_header_stub(tmp_path):
    slide = {
        "slide_number": 1,
        "render": {"slide_size_pt": {"width": 960, "height": 540}},
        "content_elements": [{
            "kind": "table", "rows": 2, "cols": 2,
            "geometry_pt": {"x_pt": 80, "y_pt": 120, "width_pt": 480, "height_pt": 160},
            "preview": [["Текст модели", "Заголовок"], ["Строка", "42"]],
            "table": {"structure": {"header_row": 0, "blank_stub_cells": [[0, 0]]}},
        }],
    }
    pptx = build_editable_pptx_from_catalog({"slides": {"slides": [slide]}}, tmp_path)
    table = next(shape.table for shape in Presentation(BytesIO(pptx)).slides[0].shapes if shape.has_table)
    assert table.cell(0, 0).text == ""
    assert table.cell(0, 1).text == "Заголовок"


def test_pptx_keeps_text_editable_with_italic_and_underline(tmp_path):
    report = {"typography": {"visibility": {"slide_size_pt": {"width": 960, "height": 540}}}}
    source = {
        "export_scene_version": 1,
        "content_elements": [{
            "kind": "text",
            "text": "Редактируемый текст",
            "geometry_pt": {"x_pt": 40, "y_pt": 40, "width_pt": 500, "height_pt": 70},
            "typography": {
                "family": "Arial", "size_pt": 24, "color": "#123456",
                "italic": True, "underline": True,
            },
        }],
    }
    scene = build_generated_export_scene(report, "Demo", [source])
    pptx = build_editable_pptx_from_catalog({"slides": {"slides": scene["slides"]}}, tmp_path)
    presentation = Presentation(BytesIO(pptx))
    run = presentation.slides[0].shapes[0].text_frame.paragraphs[0].runs[0]

    assert run.text == "Редактируемый текст"
    assert run.font.italic is True
    assert run.font.underline is True
    assert str(run.font.color.rgb) == "123456"
