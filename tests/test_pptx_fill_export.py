import json
import unittest
import zipfile
from io import BytesIO
from pathlib import Path

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Emu

from app.pptx_fill_export import apply_catalog_fill, build_grad_fill_xml, build_solid_color_fill_xml
from app.slides_pptx_builder import build_editable_pptx_from_catalog, disable_shape_effects


class PptxFillExportTest(unittest.TestCase):
    def test_build_linear_grad_fill_xml(self):
        xml = build_grad_fill_xml({
            "kind": "linear_gradient",
            "angle_deg": 90.0,
            "stops": [
                {"position": 0.0, "color": "#111111"},
                {"position": 1.0, "color": "#222222"},
            ],
        })
        self.assertIn("a:gradFill", xml or "")
        self.assertIn("a:lin", xml or "")
        self.assertIn("111111", xml or "")

    def test_build_radial_grad_fill_xml(self):
        xml = build_grad_fill_xml({
            "kind": "radial_gradient",
            "path": "circle",
            "fill_to_rect": {"l": 0.5, "t": 0.5, "r": 0.5, "b": 0.5},
            "stops": [
                {"position": 0.0, "color": "#AA0000"},
                {"position": 1.0, "color": "#00AA00"},
            ],
        })
        self.assertIn("a:path", xml or "")
        self.assertIn("fillToRect", xml or "")

    def test_apply_catalog_fill_on_shape(self):
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Emu(914400), Emu(914400))
        apply_catalog_fill(shape, {
            "kind": "linear_gradient",
            "angle_deg": 0.0,
            "stops": [
                {"position": 0.0, "color": "#FF0000"},
                {"position": 1.0, "color": "#0000FF"},
            ],
        })
        output = BytesIO()
        prs.save(output)
        xml = zipfile.ZipFile(BytesIO(output.getvalue())).read("ppt/slides/slide1.xml").decode()
        self.assertIn("gradFill", xml)
        self.assertIn("FF0000", xml.upper())

    def test_build_editable_pptx_exports_gradient_fill(self):
        root = Path("output/8b4229925584")
        if not (root / "report.json").is_file():
            self.skipTest("gradient fixture missing")
        report = json.loads((root / "report.json").read_text())
        pptx_bytes = build_editable_pptx_from_catalog(report, root / "assets")
        with zipfile.ZipFile(BytesIO(pptx_bytes)) as archive:
            slide_xml = archive.read("ppt/slides/slide21.xml").decode()
        self.assertIn("gradFill", slide_xml)

    def test_build_solid_color_fill_xml_preserves_alpha(self):
        xml = build_solid_color_fill_xml({
            "kind": "solid",
            "color": "#0077FF",
            "alpha": 0.502,
            "type": "srgb",
        })
        self.assertIn('val="0077FF"', xml)
        self.assertIn('alpha val="50200"', xml)
