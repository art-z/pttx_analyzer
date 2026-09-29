import json
import unittest
import zipfile
from io import BytesIO
from pathlib import Path

from app.slide_layer_filter import filter_slide_render_layers, should_render_slide_layer
from app.slides_export_queue import build_export_queue
from app.slides_pptx_builder import build_editable_pptx_from_catalog, slide_export_items




def test_build_export_queue_preserves_row_flex_direction():
    elements = [
        {
            "kind": "text", "text_group_id": "g1", "text_line_index": 0, "z_index": 1, "text": "Label",
            "text_group_spacing_pt": {"flex_stack_direction": "row", "flex_stack_align_items": "center", "item_gap_pt": 12},
        },
        {"kind": "text", "text_group_id": "g1", "text_line_index": 1, "z_index": 1, "text": "Value"},
    ]
    queue = build_export_queue(elements)
    assert queue[0]["kind"] == "flex_stack"
    assert queue[0]["layout"] == "row"


def test_row_flex_stack_exports_as_separate_textboxes():
    from pptx import Presentation

    from app.slides_pptx_builder import add_flex_stack_textbox

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    members = [
        {
            "kind": "text", "text": "Label", "text_group_id": "g1", "text_line_index": 0,
            "geometry_pt": {"x_pt": 40, "y_pt": 50, "width_pt": 80, "height_pt": 24},
            "text_group_spacing_pt": {"flex_stack_direction": "row", "item_gap_pt": 12},
            "typography": {"size_pt": 12},
        },
        {
            "kind": "text", "text": "Value", "text_group_id": "g1", "text_line_index": 1,
            "geometry_pt": {"x_pt": 132, "y_pt": 50, "width_pt": 120, "height_pt": 24},
            "typography": {"size_pt": 12},
        },
    ]
    add_flex_stack_textbox(slide, members, "row")
    text_shapes = [shape for shape in slide.shapes if getattr(shape, "has_text_frame", False)]
    assert len(text_shapes) == 2
    assert text_shapes[0].left < text_shapes[1].left
    assert text_shapes[0].text == "Label"
    assert text_shapes[1].text == "Value"


class SlideLayerFilterTest(unittest.TestCase):
    def test_hides_decorative_master_title_when_slide_has_content(self):
        slide = {
            "content_elements": [{"kind": "text", "placeholder_type": "title", "text": "Real title"}],
        }
        layer = {
            "kind": "text",
            "decorative": True,
            "source_scope": "master",
            "placeholder_type": "title",
            "text": "ОБРАЗЕЦ ЗАГОЛОВКА",
        }
        self.assertFalse(should_render_slide_layer(layer, slide))

    def test_keeps_background_layers(self):
        slide = {"content_elements": [{"kind": "text", "text": "Body"}]}
        bg = {"kind": "image", "source_scope": "slide_bg", "asset": "image30.png"}
        self.assertEqual(filter_slide_render_layers([bg], slide), [bg])


class EditableSlidesExportTest(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.root = Path("output/9deb32f8cd1a")
        if not (cls.root / "report.json").is_file():
            raise unittest.SkipTest("fixture report missing")
        cls.report = json.loads((cls.root / "report.json").read_text())

    def test_build_export_queue_groups_split_text(self):
        elements = [
            {"kind": "text", "text_group_id": "g1", "text_line_index": 0, "z_index": 1, "text": "Line 1"},
            {"kind": "text", "text_group_id": "g1", "text_line_index": 1, "z_index": 1, "text": "Line 2"},
            {"kind": "text", "z_index": 2, "text": "Solo"},
        ]
        queue = build_export_queue(elements)
        self.assertEqual(len(queue), 2)
        self.assertEqual(queue[0]["kind"], "flex_stack")
        self.assertEqual(len(queue[0]["members"]), 2)
        self.assertEqual(queue[1]["kind"], "element")


    def test_exported_slides_include_bullets(self):
        slide = next(item for item in self.report["slides"]["slides"] if item.get("slide_number") == 3)
        has_bullets = any(
            element.get("bullet")
            for element in slide.get("content_elements") or []
        )
        self.assertTrue(has_bullets)

        pptx_bytes = build_editable_pptx_from_catalog(self.report, self.root / "assets")
        with zipfile.ZipFile(BytesIO(pptx_bytes)) as archive:
            slide_xml = archive.read("ppt/slides/slide3.xml").decode()
        self.assertIn("buChar", slide_xml)
        self.assertIn("buClr", slide_xml)

    def test_exported_slides_have_no_shape_shadows(self):
        pptx_bytes = build_editable_pptx_from_catalog(self.report, self.root / "assets")
        with zipfile.ZipFile(BytesIO(pptx_bytes)) as archive:
            slide_xml = "".join(
                archive.read(name).decode()
                for name in archive.namelist()
                if name.startswith("ppt/slides/slide") and name.endswith(".xml")
            )
        self.assertNotIn("outerShdw", slide_xml)
        self.assertNotIn("innerShdw", slide_xml)

    def test_slide_export_items_merge_split_text_groups(self):
        slide = next(
            item for item in self.report["slides"]["slides"]
            if any(element.get("text_group_id") for element in item.get("content_elements") or [])
        )
        grouped = [item for item in slide_export_items(slide) if item.get("kind") == "flex_stack"]
        self.assertTrue(grouped)
        self.assertGreater(len(grouped[0]["members"]), 1)

    def test_disable_shape_effects_skips_graphic_frame(self):
        from pptx import Presentation
        from pptx.util import Emu

        from app.slides_pptx_builder import disable_shape_effects

        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        table_shape = slide.shapes.add_table(2, 2, 0, 0, Emu(914400), Emu(914400))
        disable_shape_effects(table_shape)

    def test_build_editable_pptx_from_catalog(self):
        pptx_bytes = build_editable_pptx_from_catalog(self.report, self.root / "assets")
        with zipfile.ZipFile(BytesIO(pptx_bytes)) as archive:
            slide_parts = [
                name for name in archive.namelist()
                if name.startswith("ppt/slides/slide") and name.endswith(".xml")
            ]
        self.assertEqual(len(slide_parts), len(self.report["slides"]["slides"]))

    def test_slide_30_exports_outline_shapes(self):
        from app.pptx import PPTXPackage
        from app.slide_catalog import extract_slide_catalog

        if not (self.root / "source.pptx").is_file():
            self.skipTest("fixture pptx missing")
        package = PPTXPackage(self.root / "source.pptx")
        catalog = extract_slide_catalog(
            package,
            self.report.get("theme") or {},
            self.report.get("slide_templates"),
        )
        slide = next(item for item in catalog["slides"] if item["slide_number"] == 30)
        line_items = [
            item["element"]
            for item in slide_export_items(slide)
            if item.get("kind") == "element" and item.get("element", {}).get("kind") == "line"
        ]
        self.assertTrue(any(item.get("outline_path") for item in line_items))
        self.assertTrue(any(item.get("line", {}).get("outline") == "rect" for item in line_items))
