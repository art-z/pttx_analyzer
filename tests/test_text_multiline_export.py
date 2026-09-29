import json
import re
import unittest
import zipfile
from io import BytesIO
from pathlib import Path

from pptx import Presentation

from app.slides_export_queue import build_export_queue
from app.slides_pptx_builder import (
    _add_text_paragraph,
    add_flex_stack_textbox,
    add_textbox,
    build_editable_pptx_from_catalog,
)


def _run_colors(xml: str) -> list[str]:
    return re.findall(
        r"<a:r>.*?<a:rPr[^>]*>.*?<a:solidFill><a:srgbClr val=\"([0-9A-F]+)\"",
        xml,
        re.S | re.I,
    )




def test_text_runs_export_language_for_pdf_word_wrapping():
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_textbox(slide, {
        "kind": "text",
        "text": "Flow: AI‑сервис для корпоративных презентаций",
        "geometry_pt": {"x_pt": 40, "y_pt": 40, "width_pt": 260, "height_pt": 80},
        "typography": {"family": "Play", "size_pt": 24, "color": "#000000"},
    })
    output = BytesIO()
    prs.save(output)
    xml = zipfile.ZipFile(BytesIO(output.getvalue())).read("ppt/slides/slide1.xml").decode()
    assert 'lang="ru-RU"' in xml
    assert 'noProof="1"' in xml
    assert 'dirty="0"' in xml
    assert '<a:cs typeface="Play"/>' in xml
    assert '<a:ea typeface="Play"/>' in xml

class TextMultilineExportTest(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.root = Path("output/3330a23478de")
        if not (cls.root / "report.json").is_file():
            cls.root = Path("output/6523b92ae85c")
        if not (cls.root / "report.json").is_file():
            raise unittest.SkipTest("multiline text fixture missing")
        cls.report = json.loads((cls.root / "report.json").read_text())

    def _slide(self, slide_number: int) -> dict:
        return next(item for item in self.report["slides"]["slides"] if item["slide_number"] == slide_number)

    def _text_element(self, slide_number: int, element_id: str) -> dict:
        slide = self._slide(slide_number)
        return next(item for item in slide["content_elements"] if item.get("element_id") == element_id)

    def test_add_text_paragraph_styles_all_soft_break_runs(self):
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        text_frame = slide.shapes.add_textbox(0, 0, 914400, 914400).text_frame
        _add_text_paragraph(
            text_frame,
            index=0,
            text="Шрифт для заголовков\nVK Sans Display Medium",
            typography={"family": "Arial", "size_pt": 12.0, "color": "#0077FF"},
            spacing=None,
        )
        output = BytesIO()
        prs.save(output)
        xml = zipfile.ZipFile(BytesIO(output.getvalue())).read("ppt/slides/slide1.xml").decode()
        colors = _run_colors(xml)
        self.assertGreaterEqual(len(colors), 2)
        self.assertTrue(all(color.upper() == "0077FF" for color in colors))
        self.assertEqual(len(re.findall(r"<a:br/>", xml)), 1)


    def test_flex_stack_member_multiline_export(self):
        slide = self._slide(8)
        queue = build_export_queue(slide["content_elements"])
        stack = next(item for item in queue if item.get("kind") == "flex_stack")
        member = stack["members"][0]
        self.assertIn("\n", member["text"])

        prs = Presentation()
        slide_obj = prs.slides.add_slide(prs.slide_layouts[6])
        add_flex_stack_textbox(slide_obj, stack["members"])
        output = BytesIO()
        prs.save(output)
        xml = zipfile.ZipFile(BytesIO(output.getvalue())).read("ppt/slides/slide1.xml").decode()
        colors = _run_colors(xml)
        expected = (member.get("typography") or {}).get("color", "").lstrip("#").upper()
        self.assertTrue(expected)
        self.assertGreaterEqual(len(colors), 2)
        self.assertTrue(all(color.upper() == expected for color in colors))

    def test_slide19_intro_text_keeps_styles_on_all_lines(self):
        slide = self._slide(19)
        queue = build_export_queue(slide["content_elements"])
        stack = queue[0]
        self.assertEqual(stack["kind"], "flex_stack")
        member = stack["members"][0]

        prs = Presentation()
        slide_obj = prs.slides.add_slide(prs.slide_layouts[6])
        add_flex_stack_textbox(slide_obj, stack["members"])
        output = BytesIO()
        prs.save(output)
        xml = zipfile.ZipFile(BytesIO(output.getvalue())).read("ppt/slides/slide1.xml").decode()
        colors = _run_colors(xml)
        expected = (member.get("typography") or {}).get("color", "").lstrip("#").upper()
        self.assertGreaterEqual(len(colors), 3)
        self.assertTrue(all(color.upper() == expected for color in colors))

    def test_slide24_second_column_multiline_export(self):
        slide = self._slide(24)
        queue = build_export_queue(slide["content_elements"])
        stacks = [item for item in queue if item.get("kind") == "flex_stack"]
        self.assertGreaterEqual(len(stacks), 2)
        member = stacks[1]["members"][0]

        prs = Presentation()
        slide_obj = prs.slides.add_slide(prs.slide_layouts[6])
        add_flex_stack_textbox(slide_obj, stacks[1]["members"])
        output = BytesIO()
        prs.save(output)
        xml = zipfile.ZipFile(BytesIO(output.getvalue())).read("ppt/slides/slide1.xml").decode()
        colors = _run_colors(xml)
        expected = (member.get("typography") or {}).get("color", "").lstrip("#").upper()
        self.assertGreaterEqual(len(colors), 2)
        self.assertTrue(all(color.upper() == expected for color in colors))

    def test_add_textbox_keeps_soft_breaks_in_single_paragraph(self):
        element = self._text_element(24, "slide24_slide_text_102")
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        add_textbox(slide, element)
        output = BytesIO()
        prs.save(output)
        xml = zipfile.ZipFile(BytesIO(output.getvalue())).read("ppt/slides/slide1.xml").decode()
        colors = _run_colors(xml)
        expected = element["typography"]["color"].lstrip("#").upper()
        self.assertEqual(len(re.findall(r"<a:p>", xml)), 1)
        self.assertEqual(len(re.findall(r"<a:br/>", xml)), 1)
        self.assertGreaterEqual(len(colors), 2)
        self.assertTrue(all(color.upper() == expected for color in colors))

    def test_build_editable_pptx_slide8_multiline_text(self):
        pptx_bytes = build_editable_pptx_from_catalog(self.report, self.root / "assets")
        with zipfile.ZipFile(BytesIO(pptx_bytes)) as archive:
            slide_xml = archive.read("ppt/slides/slide8.xml").decode()
        self.assertIn("<a:br/>", slide_xml)
        colors = _run_colors(slide_xml)
        self.assertGreaterEqual(len(colors), 4)
