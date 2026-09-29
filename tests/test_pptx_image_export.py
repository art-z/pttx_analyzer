import json
import unittest
import zipfile
from io import BytesIO
from pathlib import Path

from pptx import Presentation

from app.pptx_image_export import export_image, image_needs_shape_mask, _path_is_unit_ellipse
from app.slides_pptx_builder import build_editable_pptx_from_catalog, disable_shape_effects


class PptxImageExportTest(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.root = Path("output/da49ce6c4194")
        if not (cls.root / "report.json").is_file():
            raise unittest.SkipTest("masked image fixture missing")
        cls.report = json.loads((cls.root / "report.json").read_text())

    def _masked_image(self, slide_number: int = 3) -> dict:
        slide = next(item for item in self.report["slides"]["slides"] if item["slide_number"] == slide_number)
        return next(item for item in slide["content_elements"] if item.get("kind") == "image" and item.get("mask"))

    def test_image_needs_shape_mask(self):
        element = self._masked_image()
        self.assertTrue(image_needs_shape_mask(element))
        self.assertEqual(element["mask"]["kind"], "path")

    def test_export_image_with_path_mask(self):
        element = self._masked_image()
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        export_image(slide, element, self.root / "assets", disable_shape_effects=disable_shape_effects)
        output = BytesIO()
        prs.save(output)
        xml = zipfile.ZipFile(BytesIO(output.getvalue())).read("ppt/slides/slide1.xml").decode()
        self.assertIn("<p:pic>", xml)
        self.assertNotIn("<p:txBody>", xml)
        self.assertIn("blipFill", xml)
        self.assertTrue(
            _path_is_unit_ellipse(element["mask"]["path"])
            and 'prst="ellipse"' in xml
            or "custGeom" in xml
        )

    def test_slide3_masked_image_exports_as_picture_not_shape(self):
        root = Path("output/3330a23478de")
        if not (root / "report.json").is_file():
            root = self.root
        report = json.loads((root / "report.json").read_text())
        slide = next(item for item in report["slides"]["slides"] if item["slide_number"] == 3)
        element = next(
            item for item in slide["content_elements"]
            if item.get("kind") == "image" and item.get("mask") and item.get("crop")
        )
        prs = Presentation()
        slide_obj = prs.slides.add_slide(prs.slide_layouts[6])
        export_image(slide_obj, element, root / "assets", disable_shape_effects=disable_shape_effects)
        output = BytesIO()
        prs.save(output)
        xml = zipfile.ZipFile(BytesIO(output.getvalue())).read("ppt/slides/slide1.xml").decode()
        self.assertIn("<p:pic>", xml)
        self.assertNotIn("<p:sp><p:nvSpPr>", xml)
        self.assertIn('l="14949"', xml)
        self.assertIn('prst="ellipse"', xml)

    def test_path_mask_uses_cust_geom_for_non_circle_paths(self):
        if not (Path("output/6523b92ae85c") / "report.json").is_file():
            self.skipTest("6523 fixture missing")
        report = json.loads((Path("output/6523b92ae85c") / "report.json").read_text())
        slide = next(item for item in report["slides"]["slides"] if item["slide_number"] == 20)
        element = next(item for item in slide["content_elements"] if item.get("kind") == "image" and item.get("mask"))
        self.assertFalse(_path_is_unit_ellipse(element["mask"]["path"]))
        prs = Presentation()
        slide_obj = prs.slides.add_slide(prs.slide_layouts[6])
        export_image(slide_obj, element, Path("output/6523b92ae85c") / "assets", disable_shape_effects=disable_shape_effects)
        output = BytesIO()
        prs.save(output)
        xml = zipfile.ZipFile(BytesIO(output.getvalue())).read("ppt/slides/slide1.xml").decode()
        self.assertIn("custGeom", xml)
        self.assertIn("cubicBezTo", xml)

    def test_build_editable_pptx_exports_masked_image(self):
        pptx_bytes = build_editable_pptx_from_catalog(self.report, self.root / "assets")
        with zipfile.ZipFile(BytesIO(pptx_bytes)) as archive:
            slide_xml = archive.read("ppt/slides/slide3.xml").decode()
        self.assertIn("blipFill", slide_xml)
        self.assertIn("<p:pic>", slide_xml)
        self.assertTrue(
            "custGeom" in slide_xml
            or 'prst="ellipse"' in slide_xml
            or "roundRect" in slide_xml
        )
        masked_pic = slide_xml[slide_xml.index("<p:pic>") : slide_xml.index("</p:pic>", slide_xml.index("<p:pic>")) + len("</p:pic>")]
        self.assertNotIn("<p:txBody>", masked_pic)

    def test_export_image_applies_horizontal_flip(self):
        root = Path("output/db940ddcba3b")
        if not (root / "report.json").is_file():
            self.skipTest("db940 fixture missing")
        report = json.loads((root / "report.json").read_text())
        slide = next(item for item in report["slides"]["slides"] if item["slide_number"] == 2)
        element = next(
            layer for layer in slide["render"]["layers"]
            if layer.get("kind") == "image" and (layer.get("flip") or {}).get("h")
        )
        prs = Presentation()
        slide_obj = prs.slides.add_slide(prs.slide_layouts[6])
        export_image(slide_obj, element, root / "assets", disable_shape_effects=disable_shape_effects)
        output = BytesIO()
        prs.save(output)
        xml = zipfile.ZipFile(BytesIO(output.getvalue())).read("ppt/slides/slide1.xml").decode()
        self.assertIn('flipH="1"', xml)
        self.assertNotIn('flipV="1"', xml)

    def test_build_editable_pptx_exports_flipped_images(self):
        root = Path("output/db940ddcba3b")
        if not (root / "report.json").is_file():
            self.skipTest("db940 fixture missing")
        report = json.loads((root / "report.json").read_text())
        pptx_bytes = build_editable_pptx_from_catalog(report, root / "assets")
        with zipfile.ZipFile(BytesIO(pptx_bytes)) as archive:
            slide_xml = archive.read("ppt/slides/slide2.xml").decode()
        self.assertGreaterEqual(slide_xml.count('flipH="1"'), 1)
