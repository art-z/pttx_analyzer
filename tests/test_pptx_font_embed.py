import json
import unittest
import zipfile
from io import BytesIO
from pathlib import Path

from pptx import Presentation

from app.font_registry import collect_embeddable_fonts
from app.pptx_font_embed import embed_fonts_in_pptx
from app.slides_pptx_builder import build_editable_pptx_from_catalog


class PptxFontEmbedTest(unittest.TestCase):
    def test_collect_embeddable_fonts_from_report(self):
        root = Path("output/ab9e08a8c1f5")
        if not (root / "report.json").is_file():
            self.skipTest("fixture missing")
        report = json.loads((root / "report.json").read_text())
        fonts_root = Path(__file__).resolve().parents[1] / "app" / "fonts"
        families = collect_embeddable_fonts(report, fonts_root)
        play = next((item for item in families if item.typeface == "Play"), None)
        self.assertIsNotNone(play)
        self.assertIn("regular", play.faces)
        self.assertTrue(play.faces["regular"].path.is_file())

    def test_embed_fonts_in_pptx_writes_fntdata_parts(self):
        fonts_root = Path(__file__).resolve().parents[1] / "app" / "fonts"
        families = collect_embeddable_fonts(
            {"slides": {"slides": [{"content_elements": [{"typography": {"family": "Play", "size_pt": 14.0}}]}]}},
            fonts_root,
        )
        self.assertTrue(families)

        prs = Presentation()
        prs.slides.add_slide(prs.slide_layouts[6])
        buffer = BytesIO()
        prs.save(buffer)
        embedded = embed_fonts_in_pptx(buffer.getvalue(), families)

        with zipfile.ZipFile(BytesIO(embedded)) as archive:
            names = archive.namelist()
            self.assertTrue(any(name.startswith("ppt/fonts/font") and name.endswith(".fntdata") for name in names))
            presentation = archive.read("ppt/presentation.xml").decode()
            self.assertIn("embeddedFontLst", presentation)
            self.assertIn('typeface="Play"', presentation)
            self.assertIn('embedTrueTypeFonts="1"', presentation)
            self.assertIn('saveSubsetFonts="0"', presentation)

    def test_build_editable_pptx_embeds_play_font(self):
        root = Path("output/ab9e08a8c1f5")
        if not (root / "report.json").is_file():
            self.skipTest("fixture missing")
        report = json.loads((root / "report.json").read_text())
        pptx_bytes = build_editable_pptx_from_catalog(report, root / "assets")
        with zipfile.ZipFile(BytesIO(pptx_bytes)) as archive:
            names = archive.namelist()
            self.assertTrue(any(name.endswith(".fntdata") for name in names))
            presentation = archive.read("ppt/presentation.xml").decode()
            self.assertIn('typeface="Play"', presentation)


if __name__ == "__main__":
    unittest.main()
