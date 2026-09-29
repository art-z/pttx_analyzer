import io
import zipfile
from pathlib import Path

from pptx import Presentation

from app.slides_pptx_builder import build_editable_pptx_from_catalog


def test_generated_chart_and_diagram_export_as_editable_objects():
    render = {
        "slide_size_pt": {"width": 960, "height": 540},
        "background_color": "#FFFFFF",
    }
    report = {
        "slides": {
            "slides": [
                {
                    "slide_number": 1,
                    "render": render,
                    "content_elements": [{
                        "kind": "chart",
                        "geometry_pt": {"x_pt": 60, "y_pt": 120, "width_pt": 840, "height_pt": 340},
                        "chart_type": "bar",
                        "chart": {
                            "type": "bar",
                            "categories_preview": ["Flow", "Manual"],
                            "series": [{"name": "Hours", "values_preview": [4, 12]}],
                            "style_tokens": {
                                "series_palette": [
                                    {"color": "#0077FF"},
                                    {"color": "#FF3885"},
                                ],
                                "category_axis": {
                                    "visible": True,
                                    "typography": {"family": "Aptos", "size_pt": 11, "color": "#223344"},
                                },
                                "value_axis": {
                                    "visible": True,
                                    "typography": {"family": "Aptos", "size_pt": 9, "color": "#445566"},
                                },
                                "axis_line": {"visible": False},
                                "grid_line": {
                                    "visible": True,
                                    "horizontal_visible": True,
                                    "vertical_visible": False,
                                    "color": "#DDEEFF",
                                    "width_pt": 0.5,
                                },
                            },
                        },
                    }],
                },
                {
                    "slide_number": 2,
                    "render": render,
                    "content_elements": [{
                        "kind": "diagram",
                        "geometry_pt": {"x_pt": 60, "y_pt": 120, "width_pt": 840, "height_pt": 340},
                        "diagram_type": "flow",
                        "preview_texts": ["Upload", "Analyze", "Build"],
                        "style_tokens": {
                            "connector": {"color": {"color": "#0077FF"}, "width_pt": 1},
                            "node": {
                                "fill": {"kind": "solid", "color": "#EBF3F9"},
                                "typography": {"size_pt": 14, "color": "#000000"},
                            },
                            "arrow": {"head_type": "none", "tail_type": "triangle"},
                        },
                    }],
                },
            ],
        },
    }

    pptx_bytes = build_editable_pptx_from_catalog(report, Path("/tmp"))
    with zipfile.ZipFile(io.BytesIO(pptx_bytes)) as archive:
        assert "ppt/charts/chart1.xml" in archive.namelist()

    presentation = Presentation(io.BytesIO(pptx_bytes))
    chart = presentation.slides[0].shapes[0].chart
    assert presentation.slides[0].shapes[0].has_chart
    assert chart.category_axis.tick_labels.font.name == "Aptos"
    assert chart.category_axis.tick_labels.font.size.pt == 11
    assert str(chart.value_axis.tick_labels.font.color.rgb) == "445566"
    assert chart.value_axis.has_major_gridlines
    assert not chart.category_axis.has_major_gridlines
    assert [
        str(point.format.fill.fore_color.rgb)
        for point in chart.series[0].points
    ] == ["0077FF", "FF3885"]
    assert 'varyColors val="1"' in chart._chartSpace.xml
    diagram_shapes = presentation.slides[1].shapes
    assert len(diagram_shapes) == 5  # 3 nodes + 2 editable connectors
    assert [shape.text for shape in diagram_shapes if shape.has_text_frame][:3] == ["Upload", "Analyze", "Build"]


def test_generated_pie_export_applies_design_system_palette_per_slice():
    report = {
        "slides": {
            "slides": [{
                "slide_number": 1,
                "render": {
                    "slide_size_pt": {"width": 960, "height": 540},
                    "background_color": "#FFFFFF",
                },
                "content_elements": [{
                    "kind": "chart",
                    "geometry_pt": {"x_pt": 120, "y_pt": 80, "width_pt": 600, "height_pt": 360},
                    "chart_type": "pie",
                    "chart": {
                        "type": "pie",
                        "categories_preview": ["A", "B", "C"],
                        "series": [{"name": "Share", "values_preview": [50, 30, 20]}],
                        "style_tokens": {
                            "series_palette": [
                                {"color": "#112233"},
                                {"color": "#445566"},
                                {"color": "#778899"},
                            ],
                        },
                    },
                }],
            }],
        },
    }

    pptx_bytes = build_editable_pptx_from_catalog(report, Path("/tmp"))
    presentation = Presentation(io.BytesIO(pptx_bytes))
    points = presentation.slides[0].shapes[0].chart.series[0].points

    assert [
        str(point.format.fill.fore_color.rgb)
        for point in points
    ] == ["112233", "445566", "778899"]


def test_generated_doughnut_export_keeps_computed_hole_size():
    report = {
        "slides": {"slides": [{
            "slide_number": 1,
            "render": {"slide_size_pt": {"width": 960, "height": 540}},
            "content_elements": [{
                "kind": "chart",
                "geometry_pt": {"x_pt": 200, "y_pt": 100, "width_pt": 300, "height_pt": 300},
                "chart_type": "doughnut",
                "subtype": {"hole_size": 68},
                "chart": {
                    "type": "doughnut",
                    "subtype": {"hole_size": 68},
                    "categories_preview": ["A", "B"],
                    "series": [{"name": "Share", "values_preview": [60, 40]}],
                    "center_metric_preview": {"value": "68", "unit": "%"},
                    "style_tokens": {"circular_layout": {"center_metric": {
                        "enabled": True,
                        "preview": {"value": "68", "unit": "%"},
                        "value_typography": {"family": "Aptos", "size_pt": 24},
                        "unit_typography": {"family": "Aptos", "size_pt": 14},
                    }}},
                },
            }],
        }]},
    }
    pptx_bytes = build_editable_pptx_from_catalog(report, Path("/tmp"))
    presentation = Presentation(io.BytesIO(pptx_bytes))
    self_chart = presentation.slides[0].shapes[0].chart
    hole_size = self_chart._chartSpace.xpath(".//c:doughnutChart/c:holeSize")[0]
    assert hole_size.get("val") == "68"
    assert presentation.slides[0].shapes[1].text == "68%"


def test_generated_doughnut_keeps_square_plot_and_safe_metric_box_before_legend():
    report = {
        "slides": {"slides": [{
            "slide_number": 1,
            "render": {"slide_size_pt": {"width": 960, "height": 540}},
            "content_elements": [{
                "kind": "chart",
                "geometry_pt": {"x_pt": 45, "y_pt": 80, "width_pt": 348.62, "height_pt": 404},
                "chart_type": "doughnut",
                "subtype": {"hole_size": 64},
                "chart": {
                    "type": "doughnut",
                    "subtype": {"hole_size": 64},
                    "categories_preview": ["A", "B", "C", "D"],
                    "series": [{"name": "Share", "values_preview": [42, 32, 22, 12]}],
                    "center_metric_preview": {"value": "22", "unit": "%"},
                    "style_tokens": {
                        "legend": {"visible": True, "position": "below"},
                        "circular_layout": {"center_metric": {
                            "enabled": True,
                            "value_typography": {"family": "Play", "size_pt": 88, "bold": False},
                            "unit_typography": {"family": "Play", "size_pt": 60, "bold": False},
                        }},
                    },
                },
            }],
        }]},
    }
    pptx_bytes = build_editable_pptx_from_catalog(report, Path("/tmp"))
    presentation = Presentation(io.BytesIO(pptx_bytes))
    chart_shape, metric_shape = list(presentation.slides[0].shapes)[:2]
    assert chart_shape.width == round(348.62 * 12700)
    assert chart_shape.height == round(404 * 12700)
    expected_safe_width = chart_shape.width * 0.82 * 0.64
    assert abs(metric_shape.width - expected_safe_width) < 2
    assert metric_shape.left == chart_shape.left + (chart_shape.width - metric_shape.width) // 2
    assert metric_shape.top + metric_shape.height <= chart_shape.top + chart_shape.width
    assert all(run.font.bold is False for run in metric_shape.text_frame.paragraphs[0].runs)


def test_bar_export_prefers_design_system_fill_variants_over_deck_palette():
    report = {
        "slides": {
            "slides": [{
                "slide_number": 1,
                "render": {
                    "slide_size_pt": {"width": 960, "height": 540},
                    "background_color": "#FFFFFF",
                },
                "content_elements": [{
                    "kind": "chart",
                    "geometry_pt": {"x_pt": 60, "y_pt": 80, "width_pt": 800, "height_pt": 360},
                    "chart_type": "bar",
                    "chart": {
                        "type": "bar",
                        "categories_preview": ["A", "B", "C"],
                        "series": [{"name": "Value", "values_preview": [1, 2, 3]}],
                        "style_tokens": {
                            "series_palette": [
                                {"color": "#B3D1E8"},
                                {"color": "#262626"},
                                {"color": "#FDE53C"},
                            ],
                            "series_fill_variants": [
                                {"kind": "solid", "color": "#0077FF"},
                                {"kind": "solid", "color": "#FF3885"},
                            ],
                        },
                    },
                }],
            }],
        },
    }

    pptx_bytes = build_editable_pptx_from_catalog(report, Path("/tmp"))
    presentation = Presentation(io.BytesIO(pptx_bytes))
    chart = presentation.slides[0].shapes[0].chart

    assert [
        str(point.format.fill.fore_color.rgb)
        for point in chart.series[0].points
    ] == ["0077FF", "FF3885", "FDE53C"]
    assert 'varyColors val="1"' in chart._chartSpace.xml
