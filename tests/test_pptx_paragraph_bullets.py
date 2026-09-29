import unittest

from pptx import Presentation
from pptx.util import Emu

from app.pptx_paragraph_bullets import apply_paragraph_bullet, apply_paragraph_margins


class PptxParagraphBulletsTest(unittest.TestCase):
    def test_apply_paragraph_bullet_writes_ooxml(self):
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        shape = slide.shapes.add_textbox(0, 0, Emu(914400), Emu(914400))
        paragraph = shape.text_frame.paragraphs[0]
        paragraph.text = "Item"

        apply_paragraph_margins(
            paragraph,
            {"margin_left": 22.5, "indent": -22.5},
        )
        apply_paragraph_bullet(
            paragraph,
            {"kind": "char", "char": "•", "font": "Arial", "color": "#FD0C50", "size_pt": 16.0},
        )

        xml = paragraph._p.xml
        self.assertIn("buChar", xml)
        self.assertIn("buFont", xml)
        self.assertIn("buClr", xml)
        self.assertIn('val="FD0C50"', xml)
        self.assertIn("marL=", xml)
        self.assertIn("indent=", xml)

    def test_apply_auto_number_bullet(self):
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        shape = slide.shapes.add_textbox(0, 0, Emu(914400), Emu(914400))
        paragraph = shape.text_frame.paragraphs[0]
        apply_paragraph_bullet(
            paragraph,
            {"kind": "auto", "num_type": "arabicPeriod", "start_at": 1},
        )
        self.assertIn("buAutoNum", paragraph._p.xml)
