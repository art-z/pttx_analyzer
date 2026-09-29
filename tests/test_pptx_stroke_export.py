import json
import unittest
import zipfile
from io import BytesIO
from pathlib import Path

from pptx import Presentation
from pptx.enum.shapes import MSO_CONNECTOR

from app.pptx_stroke_export import apply_line_arrow_ends
from app.slides_pptx_builder import build_editable_pptx_from_catalog


class PptxStrokeExportTest(unittest.TestCase):
    def test_apply_line_arrow_ends_writes_head_and_tail(self):
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        connector = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, 0, 0, 914400, 914400)
        apply_line_arrow_ends(
            connector,
            {
                "head": {"type": "triangle", "width": "med", "length": "med"},
                "tail": {"type": "stealth", "width": "lg", "length": "lg"},
            },
        )
        ln = connector.line._get_or_add_ln()
        head = ln.find("{http://schemas.openxmlformats.org/drawingml/2006/main}headEnd")
        tail = ln.find("{http://schemas.openxmlformats.org/drawingml/2006/main}tailEnd")
        self.assertIsNotNone(head)
        self.assertIsNotNone(tail)
        self.assertEqual(head.get("type"), "triangle")
        self.assertEqual(head.get("w"), "med")
        self.assertEqual(head.get("len"), "med")
        self.assertEqual(tail.get("type"), "stealth")
        self.assertEqual(tail.get("w"), "lg")
        self.assertEqual(tail.get("len"), "lg")

    def test_apply_line_arrow_ends_skips_none_type(self):
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        connector = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, 0, 0, 914400, 914400)
        apply_line_arrow_ends(connector, {"head": {"type": "none"}, "tail": {"type": "triangle", "width": "lg"}})
        ln = connector.line._get_or_add_ln()
        head = ln.find("{http://schemas.openxmlformats.org/drawingml/2006/main}headEnd")
        tail = ln.find("{http://schemas.openxmlformats.org/drawingml/2006/main}tailEnd")
        self.assertIsNone(head)
        self.assertIsNotNone(tail)


class EditableSlidesArrowExportTest(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.root = Path("output/da49ce6c4194")
        if not (cls.root / "report.json").is_file():
            raise unittest.SkipTest("fixture report missing")
        cls.report = json.loads((cls.root / "report.json").read_text())

    def test_slide_13_exports_line_arrow_ends(self):
        pptx_bytes = build_editable_pptx_from_catalog(self.report, self.root / "assets")
        with zipfile.ZipFile(BytesIO(pptx_bytes)) as archive:
            slide_xml = archive.read("ppt/slides/slide13.xml").decode()
        self.assertIn("tailEnd", slide_xml)
        self.assertIn('type="triangle"', slide_xml)

    def test_slide_27_exports_stealth_line_ends(self):
        root = Path("output/d8e58f9bb017")
        if not (root / "report.json").is_file():
            self.skipTest("d8 fixture missing")
        report = json.loads((root / "report.json").read_text())
        pptx_bytes = build_editable_pptx_from_catalog(report, root / "assets")
        with zipfile.ZipFile(BytesIO(pptx_bytes)) as archive:
            slide_xml = archive.read("ppt/slides/slide27.xml").decode()
        self.assertGreaterEqual(slide_xml.count('type="stealth"'), 3)

    def test_slide_18_exports_metric_value_and_unit_sizes(self):
        pptx_bytes = build_editable_pptx_from_catalog(self.report, self.root / "assets")
        with zipfile.ZipFile(BytesIO(pptx_bytes)) as archive:
            slide_xml = archive.read("ppt/slides/slide18.xml").decode()
        self.assertIn('<a:t>23</a:t>', slide_xml)
        self.assertIn('<a:t>%</a:t>', slide_xml)
        self.assertIn('sz="8800"', slide_xml)
        self.assertIn('sz="6600"', slide_xml)
        self.assertIn('sz="23900"', slide_xml)
        self.assertIn('sz="13800"', slide_xml)
        self.assertNotIn('<a:t>23%</a:t>', slide_xml)


if __name__ == "__main__":
    unittest.main()
